#! /usr/bin/env python
# @Time    : 2025/12/22 21:47
# @Author  : afish
# @File    : DocxBase.py

import re

from docx import Document
from loguru import logger

from wordformat.agent.onnx_infer import onnx_batch_infer, onnx_single_infer
from wordformat.domains.postprocess import get_postprocess
from wordformat.settings import BATCH_SIZE
from wordformat.utils import (
    get_paragraph_numbering_text,
    para_contains_image,
    para_contains_math,
)

# ===== 目录条目识别（规则）=====
# 目录条目既不是正文也不是标题：样式为 toc 1/2/3，或文本以“制表符 + 页码”结尾
# （未套用 toc 样式、被静态化的目录）。若不排除，目录项会被当成标题参与层级
# 计数与章节状态传播，导致题注章节号和正文格式出现大量误报。
TOC_ENTRY_STYLE_PATTERN = re.compile(r"^toc\s*\d+$", re.IGNORECASE)
TOC_ENTRY_TAIL_PATTERN = re.compile(r"\t\s*(?:\d+|[ivxlcdm]+)\s*$", re.IGNORECASE)
TOC_TITLE_PATTERN = re.compile(r"^(?:目\s*录|图\s*目\s*录|表\s*目\s*录)$")


def _is_toc_entry_text(text: str) -> bool:
    return bool(TOC_ENTRY_TAIL_PATTERN.search(text))


# ===== 独立公式行识别（结构规则）=====
# 模型词表没有公式类，且 paragraph.text 不含 OMML 文字，公式段读起来往往是
# 空段或只剩右侧编号。这里按 XML 结构短路判定，不走 AI 推理。
# 编号形态：（2-1）/(2.1)/[3-12]，允许前后空白与制表符；必须整段只有编号，
# 含其他文字的段落（如"由式（2-1）可得"）仍交给模型分类，避免吞掉正文。
EQUATION_NUMBER_PATTERN = re.compile(
    r"^[\s\t]*[（(\[]\s*\d{1,2}\s*[-–—.．]\s*\d{1,3}\s*[)）\]][\s\t]*$"
)


def _is_display_formula(para, text: str) -> bool:
    """段落是否为独立公式行。

    可见文字（para.text 不含 OMML 的 m:t）只剩公式编号时判定为公式行；
    "由式（2-1）可得"这类带正文的段落仍交给模型分类。
    """
    return para_contains_math(para) and bool(EQUATION_NUMBER_PATTERN.match(text))


def _structural_item(para, text: str) -> dict:
    """按 XML 结构短路判定的段落条目（图片 / 公式 / 空段）。"""
    if para_contains_image(para):
        category, note = "figure_image", "图片段落"
    elif para_contains_math(para):
        category, note = "math_block", "公式段落"
    else:
        category, note = "body_text", "空段落"
    return {
        "category": category,
        "score": 1.0,
        "comment": note,
        "paragraph": text,
        "needs_review": False,
    }


class DocxBase:
    def __init__(self, docx_file, configpath=None, domain: str = "thesis"):
        self.re_dict = {}
        self.docx_file = docx_file
        self.domain = domain
        self.document = Document(docx_file)
        """
        以下注释掉的代码用于未来加载配置文件
        """
        # init_config(configpath)
        # try:
        #     self.config_model = get_config()  # 首次调用：触发load()
        #     self.config = self.config_model.model_dump()
        #     logger.info("配置文件验证通过")
        # except Exception as e:
        #     logger.error(f"配置加载失败: {str(e)}")
        #     raise

    def parse(self) -> list[dict]:
        # 收集所有段落（含空段）：空段按结构判定（图片/公式/空段），
        # 独立公式行短路，其余走 AI 推理
        all_paras = list(self.document.paragraphs)
        result: list = [None] * len(all_paras)
        text_indices = []
        texts_for_ai = []
        for idx, para in enumerate(all_paras):
            text = para.text.strip()
            if not text:
                result[idx] = _structural_item(para, "")
                continue
            if TOC_ENTRY_STYLE_PATTERN.match(_paragraph_style_name(para)):
                result[idx] = _toc_entry_item(text)
                continue
            if _is_display_formula(para, text):
                # 编号文字要保留，否则 jd 重建会丢掉 (2-1)
                result[idx] = _structural_item(para, text)
                continue
            numbering_text = get_paragraph_numbering_text(para)
            full_text = f"{numbering_text} {text}" if numbering_text else para.text
            texts_for_ai.append(full_text)
            text_indices.append(idx)

        # 对非空段进行批量 AI 推理（跨批次维护 PREV 上下文状态链）
        prev_label = None
        for i in range(0, len(texts_for_ai), BATCH_SIZE):
            batch_texts = texts_for_ai[i : i + BATCH_SIZE]
            batch_indices = text_indices[i : i + BATCH_SIZE]

            try:
                batch_results, prev_label = onnx_batch_infer(batch_texts, prev_label)
            except Exception as e:
                logger.error(f"批量推理失败，降级到单条处理: {e}")
                batch_results = []
                for t in batch_texts:
                    r = onnx_single_infer(t, prev_label)
                    batch_results.append(r)
                    if r["label"]:
                        prev_label = r["label"]

            for idx, text, pred in zip(
                batch_indices, batch_texts, batch_results, strict=False
            ):
                response = {
                    "category": pred["label"],
                    "score": pred["score"],
                    "comment": f"置信度：{pred['score']:.4f}",
                    "paragraph": text,
                }
                result[idx] = response

        assert all(r is not None for r in result), "存在未处理的段落"

        # 目录条目是结构信息（既非正文也非标题），先按样式/文本排除，
        # 避免后续领域规则把它们当成标题参与编号与状态传播。
        _fix_toc_entries(result, all_paras)

        # 领域分类后处理：阈值应用 + 领域规则链由领域模块整链执行并注册到
        # 框架注册表（domains/postprocess.py）。领域未装配（未注册）时跳过，
        # 顺序约束与阈值常量随论文规则位于 domains/thesis/postprocess.py
        postprocessor = get_postprocess(self.domain)
        if postprocessor is not None:
            postprocessor(result)
        return self._interleave_tables(result)

    def _interleave_tables(self, result: list[dict]) -> list[dict]:
        """把文档中的表格按文档序导出为 table_object 条目。

        ``document.paragraphs`` 只含 body 直接子级段落，表格单元格不在其中，
        原表格在 gj JSON 中完全不可见。此处按 body XML 顺序遍历 w:p / w:tbl，
        在每个表格前记录已出现的段落数，把表格条目插入对应位置。

        表格条目为**独立项**：不占用段落槽位，带 ``exists=True`` 供 zip 侧
        跳过段落绑定（建树阶段不绑定原表格）；``table_rows`` 为全部单元格
        文本（行优先，合并单元格按 python-docx 语义拍平）。
        """
        from docx.oxml.ns import qn

        body = self.document.element.body
        tables = []
        p_count = 0
        for el in body.iterchildren():
            if el.tag == qn("w:p"):
                p_count += 1
            elif el.tag == qn("w:tbl"):
                tables.append((p_count, el))
        if not tables:
            return result

        out = list(result)
        for k, (pos, tbl_el) in enumerate(tables):
            out.insert(
                min(pos + k, len(out)),
                {
                    "category": "table_object",
                    "paragraph": "",
                    "score": 1.0,
                    "comment": "原文档表格（内容只读，不参与格式检查）",
                    "table_rows": self._extract_table_rows(tbl_el),
                    "exists": True,
                    "needs_review": False,
                },
            )
        return out

    def _extract_table_rows(self, tbl_el) -> list[list[str]]:
        """提取表格全部单元格文本（行优先）。"""
        from docx.table import Table

        table = Table(tbl_el, self.document._body)
        return [[cell.text for cell in row.cells] for row in table.rows]


def _paragraph_style_name(para) -> str:
    style = getattr(para, "style", None)
    return (getattr(style, "name", "") or "").strip()


def _toc_entry_item(text: str) -> dict:
    return {
        "category": "toc_entry",
        "score": 1.0,
        "comment": "目录条目（规则识别，跳过格式化）",
        "paragraph": text,
        "needs_review": False,
    }


def _fix_toc_entries(result: list[dict], paragraphs) -> None:
    """把目录条目标记为 ``toc_entry``，与正文/标题彻底隔离。

    两类来源：
    1. 段落样式为 ``toc 1`` / ``toc 2`` / ``toc 3``（Word 自动目录）；
    2. 静态化目录：``目录`` 标题之后、以“制表符 + 页码”结尾的连续段落。
       目录块遇到第一个非条目段落即结束，避免误伤正文。
    """
    in_toc = False
    for item, para in zip(result, paragraphs, strict=False):
        text = (para.text or "").strip()
        if TOC_ENTRY_STYLE_PATTERN.match(_paragraph_style_name(para)):
            item.update(_toc_entry_item(text))
            in_toc = True
            continue
        if not text:
            continue
        if TOC_TITLE_PATTERN.match(text):
            in_toc = True
            continue
        if in_toc and _is_toc_entry_text(text):
            item.update(_toc_entry_item(text))
            continue
        in_toc = False
