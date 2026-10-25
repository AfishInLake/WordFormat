#! /usr/bin/env python
# @Time    : 2026/9/28
# @Author  : afish
# @File    : header_footer.py
"""页眉页脚格式识别与修正（issue#93）。

Word 的页眉页脚存在于各节（section）的 header/footer part 中，不属于正文
段落流，模型分类与 FormatNode 树均不覆盖。本模块通过 hook 机制接入流水线
（见 handlers.py 中注册的 builtin.header_footer 回调，触发于
before_document_save 事件）：

- check 模式：检查每节页眉/页脚的字号、中文字体、英文字体、对齐方式及
  内容缺失，问题作为批注锚定在文档第一段（页眉/页脚 run 无法挂批注——
  python-docx 限制 comments 只能出现在 document part）；
- apply 模式：按配置修正格式，残留差异仍以批注报告。

配置段（YAML 可选，缺省用 DEFAULT_HF_CONFIG）::

    header_footer:
      enabled: true
      update_fields_on_open: false   # 打开文档时更新域（settings.xml w:updateFields）
      header:
        chinese_font_name: '宋体'
        english_font_name: 'Times New Roman'
        font_size: '小五'
        alignment: '居中对齐'
        require_content: false
        page_number: false           # 要求存在 PAGE 页码域
      footer: ...
      page_setup:            # 页边距 / 距边界 / 纸张 / 页码 / 分节开关
        top_margin: '2.54cm'
        bottom_margin: '2.54cm'
        left_margin: '3.18cm'
        right_margin: '3.18cm'
        header_distance: '1.5cm'
        footer_distance: '1.75cm'
        page_size: 'A4'              # A3 / A4 / A5 / B5 / Letter
        orientation: '纵向'          # 纵向 / 横向
        page_number_start: 1         # w:pgNumType/@w:start
        page_number_format: '小写罗马数字'
        different_first_page: false  # 首页不同（w:titlePg）
        odd_even_pages: false        # 奇偶页不同（w:evenAndOddHeaders）

页眉/页脚仅检查四项属性：字号、中文字体（eastAsia）、英文字体（ascii）、对齐。
不检查加粗/颜色/行距等易误报属性，避免对页码域等特殊内容误伤。

页边距（page_setup）作用于每个 section 的 pgMar：上/下/左/右边距与页眉/页脚距
边界，单位支持 cm / mm / in / pt（缺省按 cm）；逐节确定性校验/修正，不依赖模型分类。

纸张/页码设置写在 sectPr 上（pgSz / pgNumType / titlePg），逐节校验/修正；
奇偶页不同与"打开时更新域"是文档级开关，写在 settings.xml，只处理一次。

页码域（header/footer 的 page_number）：check 模式只报告缺失，apply 模式插入
``w:fldSimple w:instr=" PAGE "``。域结果需要 Word 刷新才显示真实页码，可配合
update_fields_on_open 打开时自动更新。

首页/偶页页眉页脚变体只在该节（或文档）已开启对应开关、且变体自带部件时才参与
检查——未开启时 Word 不显示它，检查只会误报；部件不存在表示沿用其他节。
读取部件前先判断 ``_has_definition``：python-docx 访问 .paragraphs 会给"沿用上一节"
的节凭空创建部件，check 模式不应因此改动文档。
"""

from __future__ import annotations

import copy
import re
from typing import Iterator, List, Optional, Tuple

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from loguru import logger

from wordformat.style.defs import Alignment, FontSize
from wordformat.style.reader import (
    paragraph_get_alignment,
    run_get_font_name,
    run_get_font_name_en,
    run_get_font_size_pt_or_none,
)
from wordformat.utils import has_chinese

# 页眉/页脚缺省规范（本科论文常见：小五号宋体、Times New Roman、居中）
DEFAULT_HF_CONFIG = {
    "enabled": True,
    "header": {
        "chinese_font_name": "宋体",
        "english_font_name": "Times New Roman",
        "font_size": "小五",  # 9pt
        "alignment": "居中对齐",
        "require_content": False,
    },
    "footer": {
        "chinese_font_name": "宋体",
        "english_font_name": "Times New Roman",
        "font_size": "小五",  # 9pt
        "alignment": "居中对齐",
        "require_content": False,
    },
    # page_setup（页边距/页眉页脚距边界）为 opt-in：不放进缺省配置，
    # 避免对每个文档误报（各校边距差异大）；仅当 YAML 写了 header_footer.page_setup 才校验。
}

# 页边距推荐缺省（仅供文档/测试参考，不主动注入 resolve_hf_config）：
# 上/下 2.54cm、左/右 3.18cm、页眉距边界 1.5cm、页脚距边界 1.75cm（Word 默认近似）

# 页边距字段：(section 属性名 = 配置键, 中文标签)
_PAGE_FIELDS = [
    ("top_margin", "上边距"),
    ("bottom_margin", "下边距"),
    ("left_margin", "左边距"),
    ("right_margin", "右边距"),
    ("header_distance", "页眉距边界"),
    ("footer_distance", "页脚距边界"),
]

# 常见纸张的纵向尺寸（宽, 高）单位 cm；"横向"时互换
_PAGE_SIZES = {
    "A3": (29.7, 42.0),
    "A4": (21.0, 29.7),
    "A5": (14.8, 21.0),
    "B5": (17.6, 25.0),
    "LETTER": (21.59, 27.94),
}
# pgSz 存的是 twips，换算回 cm 有 ±0.01cm 舍入，纸张尺寸比对放宽到 0.1cm
_PAGE_SIZE_TOL_CM = 0.1
_LANDSCAPE_LABELS = {"横向", "landscape"}

# 页码格式标签 → w:pgNumType/@w:fmt；配置直接写 OOXML 值也接受
_PAGE_NUMBER_FORMATS = {
    "阿拉伯数字": "decimal",
    "阿拉伯数字补零": "decimalZero",
    "大写罗马数字": "upperRoman",
    "小写罗马数字": "lowerRoman",
    "大写英文字母": "upperLetter",
    "小写英文字母": "lowerLetter",
    "中文数字": "chineseCounting",
}
# OOXML 原值 → 原值（小写索引，允许配置直接写 decimal / upperRoman 等）
_OOXML_PAGE_FORMATS = {value.lower(): value for value in _PAGE_NUMBER_FORMATS.values()}
_OOXML_PAGE_FORMATS["none"] = "none"

# w:pgNumType 在 CT_SectPr 中的后继元素（插入时必须排在其前）
_PG_NUM_TYPE_SUCC = (
    "w:cols",
    "w:formProt",
    "w:vAlign",
    "w:noEndnote",
    "w:titlePg",
    "w:textDirection",
    "w:bidi",
    "w:rtl",
    "w:docGrid",
    "w:printerSettings",
    "w:sectPrChange",
)

# settings.xml 中两个开关的后继元素（取自 docx.oxml.settings.CT_Settings._tag_seq）
_SETTING_SUCC = {
    "w:evenAndOddHeaders": (
        "w:bookFoldRevPrinting",
        "w:bookFoldPrinting",
        "w:drawingGridHorizontalSpacing",
        "w:characterSpacingControl",
        "w:savePreviewPicture",
        "w:updateFields",
        "w:footnotePr",
        "w:endnotePr",
        "w:compat",
        "w:rsids",
        "m:mathPr",
        "w:shapeDefaults",
        "w:decimalSymbol",
        "w:listSeparator",
    ),
    "w:updateFields": (
        "w:hdrShapeDefaults",
        "w:footnotePr",
        "w:endnotePr",
        "w:compat",
        "w:docVars",
        "w:rsids",
        "m:mathPr",
        "w:themeFontLang",
        "w:clrSchemeMapping",
        "w:shapeDefaults",
        "w:decimalSymbol",
        "w:listSeparator",
    ),
}

# 页码域指令：PAGE 但排除 NUMPAGES（总页数不是页码）
_PAGE_INSTR_PATTERN = re.compile(r"(?<!\w)PAGE(?!\w)", re.IGNORECASE)

# 页眉页脚字号容差：小五(9pt) / 五号(10.5pt) 在页眉里视觉等价，各校常混用；
# 只有明显大于该档（如正文小四 12pt）才视为错误。
_HF_SIZE_TOL_PT = 1.5
_ASCII_LETTER_PATTERN = re.compile(r"[A-Za-z]")

# 页眉/页脚变体：配置键 → (中文标签, {变体名: section 属性名})
_HF_PARTS = {
    "header": (
        "页眉",
        {
            "default": "header",
            "first": "first_page_header",
            "even": "even_page_header",
        },
    ),
    "footer": (
        "页脚",
        {
            "default": "footer",
            "first": "first_page_footer",
            "even": "even_page_footer",
        },
    ),
}
_HF_VARIANT_LABELS = {"first": "首页", "even": "偶页"}


# 文档级（settings.xml）开关：(配置键, 中文标签, OOXML 标签)
_DOC_FLAGS = (
    ("odd_even_pages", "奇偶页不同", "w:evenAndOddHeaders"),
    ("update_fields_on_open", "打开时更新域", "w:updateFields"),
)


def resolve_hf_config(config_model=None) -> dict:
    """从配置模型提取 header_footer 段，与缺省规范深合并（用户配置优先）。

    config_model 可为 NodeConfigRoot / dict / None；禁用（enabled=false）
    返回空 dict，调用方应跳过处理。
    """
    if config_model is None:
        merged = copy.deepcopy(DEFAULT_HF_CONFIG)
    else:
        hf = (
            config_model.get("header_footer")
            if isinstance(config_model, dict)
            else getattr(config_model, "header_footer", None)
        )
        if not isinstance(hf, dict):
            hf = {}
        from wordformat.config.dotdict import deep_merge

        merged = deep_merge(copy.deepcopy(DEFAULT_HF_CONFIG), hf)
    return merged if merged.get("enabled", True) else {}


def _part_label(kind: str, sec_idx: int, variant: str = "") -> str:
    """节标签：页眉(第1节) / 页眉(第1节·首页)。"""
    suffix = f"·{variant}" if variant else ""
    return f"{kind}(第{sec_idx}节{suffix})"


def _resolve_alignment_value(label):
    """对齐标签 → WD_ALIGN_PARAGRAPH 枚举值；解析失败返回 None。"""
    try:
        return Alignment(label).rel_value
    except (ValueError, TypeError):
        return None


def _resolve_font_size_pt(label) -> Optional[float]:
    """字号标签 → pt；支持 '小五' / 12 / '12pt' 等。"""
    if label is None:
        return None
    try:
        if isinstance(label, str):
            return float(FontSize(label).rel_value)
        return float(label)
    except (ValueError, TypeError):
        return None


_CM_PER = {"cm": 1.0, "毫米": 0.1, "mm": 0.1, "英寸": 2.54, "in": 2.54, "pt": 0.035278}
_UNIT_ALIAS = {"厘米": "cm", "毫米": "mm", "inch": "in", "英寸": "in", "磅": "pt"}


def _resolve_length_cm(label) -> Optional[float]:
    """长度标签 → cm 浮点；支持 '2.54cm' / '31.8mm' / '1.5' (缺省 cm) / 数字。

    无法解析返回 None（调用方跳过该项，避免误报）。
    """
    if label is None:
        return None
    if isinstance(label, (int, float)):
        return float(label)
    pattern = r"^\s*([0-9]+(?:\.[0-9]+)?)\s*(cm|厘米|毫米|mm|英寸|in|inch|pt|磅)?\s*$"
    m = re.match(pattern, str(label).lower())
    if not m:
        return None
    unit = m.group(2) or "cm"
    unit = _UNIT_ALIAS.get(unit, unit)
    return float(m.group(1)) * _CM_PER.get(unit, 1.0)


def _resolve_page_format(label) -> Optional[str]:
    """页码格式标签 → w:pgNumType/@w:fmt；支持中文标签与 OOXML 原值。"""
    if label is None:
        return None
    text = str(label).strip()
    mapped = _PAGE_NUMBER_FORMATS.get(text)
    if mapped:
        return mapped
    return _OOXML_PAGE_FORMATS.get(text.lower())


def _resolve_page_start(label) -> Optional[int]:
    """起始页码 → int；未配置或无法解析返回 None。"""
    if label is None:
        return None
    try:
        return int(str(label).strip())
    except (TypeError, ValueError):
        return None


def _is_landscape(label) -> bool:
    return label is not None and str(label).strip().lower() in _LANDSCAPE_LABELS


def _expected_page_dims(cfg: dict) -> Optional[Tuple[float, float]]:
    """(宽cm, 高cm)：由 page_size + orientation 推出；未配置 page_size 返回 None。"""
    label = cfg.get("page_size")
    if label is None:
        return None
    dims = _PAGE_SIZES.get(str(label).strip().upper())
    if dims is None:
        return None
    return (dims[1], dims[0]) if _is_landscape(cfg.get("orientation")) else dims


def _page_orient_value(section) -> str:
    """sectPr/pgSz/@w:orient 当前值；缺省按 portrait。"""
    pg_sz = section._sectPr.find(qn("w:pgSz"))
    if pg_sz is None:
        return "portrait"
    return (pg_sz.get(qn("w:orient")) or "portrait").lower()


def _check_page_geometry(section, pos: str, cfg: dict, issues: List[str]):
    """纸张尺寸与横竖方向（pgSz）。"""
    expected = _expected_page_dims(cfg)
    width, height = section.page_width, section.page_height
    if expected and width is not None and height is not None:
        if (
            abs(width.cm - expected[0]) > _PAGE_SIZE_TOL_CM
            or abs(height.cm - expected[1]) > _PAGE_SIZE_TOL_CM
        ):
            issues.append(
                f"{pos}-纸张尺寸错误：当前{width.cm:.2f}×{height.cm:.2f}cm，"
                f"规范:{cfg.get('page_size')}"
            )
    orientation = cfg.get("orientation")
    if orientation is None:
        return
    actual = _page_orient_value(section) == "landscape"
    if actual != _is_landscape(orientation):
        issues.append(
            f"{pos}-纸张方向错误：当前{'横向' if actual else '纵向'}，规范:{orientation}"
        )


def _apply_page_geometry(section, cfg: dict):
    """写入 pgSz 宽高与 @w:orient；只配方向时对调现有宽高。"""
    from docx.shared import Cm

    orientation = cfg.get("orientation")
    expected = _expected_page_dims(cfg)
    if expected is None and orientation is not None:
        if section.page_width is None or section.page_height is None:
            return
        a, b = section.page_width.cm, section.page_height.cm
        low, high = (a, b) if a <= b else (b, a)
        expected = (high, low) if _is_landscape(orientation) else (low, high)
    if expected is None:
        return
    section.page_width = Cm(round(expected[0], 4))
    section.page_height = Cm(round(expected[1], 4))
    if orientation is not None:
        value = "landscape" if _is_landscape(orientation) else "portrait"
        section._sectPr.get_or_add_pgSz().set(qn("w:orient"), value)


def _pg_num_type(section, create: bool = False):
    """sectPr/pgNumType 元素；create=False 时只读，可能为 None。"""
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:pgNumType"))
    if node is None and create:
        node = OxmlElement("w:pgNumType")
        sect_pr.insert_element_before(node, *_PG_NUM_TYPE_SUCC)
    return node


def _check_page_number_type(section, pos: str, cfg: dict, issues: List[str]):
    """起始页码与页码格式（pgNumType）。"""
    expected_start = _resolve_page_start(cfg.get("page_number_start"))
    expected_fmt = _resolve_page_format(cfg.get("page_number_format"))
    if expected_start is None and not expected_fmt:
        return
    node = _pg_num_type(section)
    actual_start = node.get(qn("w:start")) if node is not None else None
    actual_fmt = node.get(qn("w:fmt")) if node is not None else None
    if expected_start is not None and actual_start != str(expected_start):
        issues.append(
            f"{pos}-起始页码错误：当前{actual_start or '未设置（沿用上一节）'}，"
            f"规范:{cfg.get('page_number_start')}"
        )
    if expected_fmt and actual_fmt != expected_fmt:
        issues.append(
            f"{pos}-页码格式错误：当前{actual_fmt or 'decimal'}，"
            f"规范:{cfg.get('page_number_format')}"
        )


def _apply_page_number_type(section, cfg: dict):
    expected_start = _resolve_page_start(cfg.get("page_number_start"))
    expected_fmt = _resolve_page_format(cfg.get("page_number_format"))
    if expected_start is None and not expected_fmt:
        return
    node = _pg_num_type(section, create=True)
    if expected_start is not None:
        node.set(qn("w:start"), str(expected_start))
    if expected_fmt:
        node.set(qn("w:fmt"), expected_fmt)


def _check_first_page_flag(section, pos: str, cfg: dict, issues: List[str]):
    expected = cfg.get("different_first_page")
    if expected is None:
        return
    actual = bool(getattr(section, "different_first_page_header_footer", False))
    if actual != bool(expected):
        issues.append(
            f"{pos}-首页不同页眉页脚错误：当前{'开启' if actual else '关闭'}，"
            f"规范:{'开启' if expected else '关闭'}"
        )


def _apply_first_page_flag(section, cfg: dict):
    expected = cfg.get("different_first_page")
    if expected is None:
        return
    section.different_first_page_header_footer = bool(expected)


def _settings_element(document):
    """settings.xml 根元素；部件缺失返回 None。"""
    try:
        return document.settings.element
    except (AttributeError, KeyError):
        return None


def _setting_on(document, tag: str) -> bool:
    """settings.xml 开关是否为"开"（元素缺失或 val=false 都算关）。"""
    settings = _settings_element(document)
    if settings is None:
        return False
    element = settings.find(qn(tag))
    if element is None:
        return False
    value = (element.get(qn("w:val")) or "true").lower()
    return value not in ("false", "0", "off")


def _set_setting(document, tag: str, on: bool):
    """写 settings.xml 开关：关闭时删除元素（OOXML 缺省即关）。"""
    settings = _settings_element(document)
    if settings is None:
        return
    element = settings.find(qn(tag))
    if not on:
        if element is not None:
            settings.remove(element)
        return
    if element is None:
        element = OxmlElement(tag)
        settings.insert_element_before(element, *_SETTING_SUCC.get(tag, ()))
    element.set(qn("w:val"), "true")


def _check_doc_flags(document, hf_cfg: dict, issues: List[str]):
    """文档级开关（settings.xml）：配置了才比对。"""
    for key, label, tag in _DOC_FLAGS:
        expected = hf_cfg.get(key)
        if expected is None:
            continue
        actual = _setting_on(document, tag)
        if actual != bool(expected):
            issues.append(
                f"文档-{label}错误：当前{'开启' if actual else '关闭'}，"
                f"规范:{'开启' if expected else '关闭'}"
            )


def _apply_doc_flags(document, hf_cfg: dict, remaining: List[str]):
    for key, _label, tag in _DOC_FLAGS:
        if hf_cfg.get(key) is not None:
            _set_setting(document, tag, bool(hf_cfg[key]))
    _check_doc_flags(document, hf_cfg, remaining)


def _check_page_setup(section, sec_idx: int, cfg: dict, issues: List[str]):
    """校验单节页边距、纸张、页码设置与首页不同开关，偏差 > 0.02cm 报批注。"""
    pos = f"页面(第{sec_idx}节)"
    for attr, name in _PAGE_FIELDS:
        expected = _resolve_length_cm(cfg.get(attr))
        if expected is None:
            continue
        actual_len = getattr(section, attr, None)
        if actual_len is None:
            continue
        actual = actual_len.cm
        if abs(actual - expected) > 0.02:
            issues.append(f"{pos}-{name}错误：当前{actual:.2f}cm，规范:{cfg.get(attr)}")
    _check_page_geometry(section, pos, cfg, issues)
    _check_page_number_type(section, pos, cfg, issues)
    _check_first_page_flag(section, pos, cfg, issues)


def _apply_page_setup(section, sec_idx: int, cfg: dict, remaining: List[str]):
    """修正单节页面设置；修正后复检，残留写入 remaining。"""
    from docx.shared import Cm

    for attr, _name in _PAGE_FIELDS:
        expected = _resolve_length_cm(cfg.get(attr))
        if expected is None:
            continue
        setattr(section, attr, Cm(round(expected, 4)))
    _apply_page_geometry(section, cfg)
    _apply_page_number_type(section, cfg)
    _apply_first_page_flag(section, cfg)
    _check_page_setup(section, sec_idx, cfg, remaining)


def _para_mark_font(paragraph, attr: str) -> str | None:
    """段落标记 rPr 中的字体（WPS 常把页眉字体写在 w:pPr/w:rPr）。"""
    ppr = paragraph._p.find(qn("w:pPr"))
    mark = ppr.find(qn("w:rPr")) if ppr is not None else None
    rfonts = mark.find(qn("w:rFonts")) if mark is not None else None
    if rfonts is None:
        return None
    return rfonts.get(qn(f"w:{attr}"))


def _check_paragraph(
    paragraph, kind: str, sec_idx: int, cfg: dict, issues: List[str], variant: str = ""
):
    """检查单个页眉/页脚段落（非空）：对齐 + 字号 + 中英文字体。"""
    pos = _part_label(kind, sec_idx, variant)

    # 对齐（有效值沿继承链解析，None 表示未设置 → 跳过避免误报）
    expected_alignment = _resolve_alignment_value(cfg.get("alignment"))
    if expected_alignment is not None:
        actual = paragraph_get_alignment(paragraph)
        if actual is not None and actual != expected_alignment:
            issues.append(
                f"{pos}-对齐错误：当前非{cfg.get('alignment')}，规范:{cfg.get('alignment')}"
            )

    # 字符属性：字号 + 中/英文字体（每段各检查一次即可，取首个非空 run）
    expected_size = _resolve_font_size_pt(cfg.get("font_size"))
    expected_cn = (cfg.get("chinese_font_name") or "").lower()
    expected_en = (cfg.get("english_font_name") or "").lower()
    if not (expected_size or expected_cn or expected_en):
        return

    checked_cn = checked_en = checked_size = False
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        if expected_size and not checked_size:
            # 字号沿继承链均未设置时不判定（域代码结果 run 常见，避免误报 12pt）
            actual_size = run_get_font_size_pt_or_none(run)
            if actual_size is not None:
                size_ok = abs(actual_size - expected_size) <= _HF_SIZE_TOL_PT
                if not size_ok:
                    issues.append(
                        f"{pos}-字号错误：当前{actual_size:g}pt，规范:{cfg.get('font_size')}"
                    )
            checked_size = True
        if expected_cn and not checked_cn and has_chinese(run.text):
            actual = run_get_font_name(run) or _para_mark_font(paragraph, "eastAsia")
            if actual is not None and actual.lower() != expected_cn:
                issues.append(
                    f"{pos}-中文字体错误：当前{actual}，"
                    f"规范:{cfg.get('chinese_font_name')}"
                )
            checked_cn = True
        if expected_en and not checked_en and _ASCII_LETTER_PATTERN.search(run.text):
            actual = run_get_font_name_en(run) or _para_mark_font(paragraph, "ascii")
            if actual is not None and actual.lower() != expected_en:
                issues.append(
                    f"{pos}-英文字体错误：当前{actual}，"
                    f"规范:{cfg.get('english_font_name')}"
                )
            checked_en = True
        if checked_cn and checked_en and checked_size:
            break


def _has_own_part(part) -> bool:
    """该节是否自带页眉/页脚部件（只读判断，不会创建部件）。"""
    return bool(getattr(part, "_has_definition", False))


def _owns_content(part, sec_idx: int, variant: str) -> bool:
    """本节是否自己负责该变体的内容。

    部件不存在时：第 2 节起沿用上一节、未创建的首页/偶页变体由 Word 隐藏，
    都不该报缺失；只有第 1 节的默认页眉/页脚为空才是真的空。
    """
    return _has_own_part(part) or (sec_idx == 1 and not variant)


def _content_paragraphs(part):
    """部件中的非空段落；部件不存在时返回空列表（不触发创建）。"""
    if not _has_own_part(part):
        return []
    return [p for p in part.paragraphs if p.text.strip()]


def _field_instructions(part) -> str:
    """部件内所有域指令拼成一句。

    页码域有两种写法：w:fldSimple/@w:instr 与复杂域的 w:instrText。复杂域的
    instrText 会被 Word 拆成多段（" P" + "AGE " + "\\* MERGEFORMAT"），必须直接
    相接而不能加分隔符，否则拼不出 PAGE。
    """
    simple = []
    pieces = []
    for element in part._element.xpath(".//w:fldSimple | .//w:instrText"):
        if element.tag == qn("w:fldSimple"):
            simple.append(element.get(qn("w:instr")) or "")
        else:
            pieces.append(element.text or "")
    return " ".join([*simple, "".join(pieces)])


def _has_page_field(part) -> bool:
    """部件是否含 PAGE 页码域（NUMPAGES 等总页数域不算）。"""
    return bool(_PAGE_INSTR_PATTERN.search(_field_instructions(part)))


def _insert_page_field(part, cfg: dict) -> None:
    """在部件首段末尾插入 PAGE 页码域，并按配置补齐段落对齐与字体。

    部件不存在时访问 .paragraphs 会创建它——apply 模式需要该副作用。
    """
    from docx.text.run import Run

    paragraph = part.paragraphs[0]
    expected_alignment = _resolve_alignment_value(cfg.get("alignment"))
    if expected_alignment is not None:
        paragraph.alignment = expected_alignment
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r_element = OxmlElement("w:r")
    t_element = OxmlElement("w:t")
    t_element.text = "1"  # 域缓存结果占位，Word 更新域后替换为真实页码
    r_element.append(t_element)
    fld.append(r_element)
    paragraph._p.append(fld)
    _apply_run_format(Run(r_element, paragraph), cfg)


def _page_number_issue(kind: str, sec_idx: int, variant: str) -> str:
    return (
        f"{_part_label(kind, sec_idx, variant)}-页码缺失：当前无页码域，"
        f"规范：应插入 PAGE 页码域"
    )


def _missing_content_issue(kind: str, sec_idx: int, variant: str) -> str:
    return (
        f"{_part_label(kind, sec_idx, variant)}-内容缺失：当前为空，"
        f"规范：应包含{kind}内容"
    )


def _check_missing_content(
    kind: str, sec_idx: int, variant: str, cfg: dict, sink: List[str]
):
    """空部件报告：显式 require_content 才算问题；页码域本身就是内容。"""
    if cfg.get("require_content", False) and not cfg.get("page_number"):
        sink.append(_missing_content_issue(kind, sec_idx, variant))


def _check_part(
    part, kind: str, sec_idx: int, cfg: dict, issues: List[str], variant: str = ""
):
    """检查某节的一个页眉/页脚变体。"""
    if not _owns_content(part, sec_idx, variant):
        return
    has_part = _has_own_part(part)
    if cfg.get("page_number") and not (has_part and _has_page_field(part)):
        issues.append(_page_number_issue(kind, sec_idx, variant))
    paragraphs = _content_paragraphs(part)
    if not paragraphs:
        _check_missing_content(kind, sec_idx, variant, cfg, issues)
        return
    for p in paragraphs:
        _check_paragraph(p, kind, sec_idx, cfg, issues, variant)


def _section_parts(
    section, sec_idx: int, hf_cfg: dict, even_pages_on: bool
) -> Iterator[Tuple[str, dict, object, str]]:
    """该节在 Word 中实际生效的 (类别, 子配置, 部件, 变体标签)。

    首页变体只在该节开了"首页不同"、偶页变体只在文档开了"奇偶页不同"时才显示，
    未生效的部件不参与检查。
    """
    first_on = bool(getattr(section, "different_first_page_header_footer", False))
    variants = ["default"]
    if first_on:
        variants.append("first")
    if even_pages_on:
        variants.append("even")
    for cfg_key, (kind, attrs) in _HF_PARTS.items():
        cfg = hf_cfg.get(cfg_key) or {}
        for name in variants:
            yield (
                kind,
                cfg,
                getattr(section, attrs[name]),
                _HF_VARIANT_LABELS.get(name, ""),
            )


def check_header_footer(document, config_model=None) -> List[str]:
    """检查文档所有节的页眉页脚与页面设置，返回标准批注文案列表。

    Args:
        document: python-docx Document 对象
        config_model: 配置模型（NodeConfigRoot/dict），None 用缺省规范

    Returns:
        问题文案列表（空表示无问题或未启用）
    """
    hf_config = resolve_hf_config(config_model)
    if not hf_config:
        return []
    issues: List[str] = []
    page_cfg = hf_config.get("page_setup") or {}
    even_pages_on = _setting_on(document, "w:evenAndOddHeaders")
    for idx, section in enumerate(document.sections, start=1):
        for kind, cfg, part, variant in _section_parts(
            section, idx, hf_config, even_pages_on
        ):
            _check_part(part, kind, idx, cfg, issues, variant)
        if page_cfg:
            _check_page_setup(section, idx, page_cfg, issues)
    _check_doc_flags(document, hf_config, issues)
    logger.debug(f"页眉页脚/页面设置检查完成：{len(issues)} 个问题")
    return issues


def _apply_run_format(run, cfg: dict) -> None:
    """修正单个 run 的字号与中/英文字体。"""
    if not run.text.strip():
        return
    from docx.shared import Pt

    expected_size = _resolve_font_size_pt(cfg.get("font_size"))
    if expected_size:
        run.font.size = Pt(expected_size)

    cn_name = cfg.get("chinese_font_name")
    en_name = cfg.get("english_font_name")
    if cn_name or en_name:
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        if cn_name:
            rFonts.set(qn("w:eastAsia"), cn_name)
        if en_name:
            rFonts.set(qn("w:ascii"), en_name)
            rFonts.set(qn("w:hAnsi"), en_name)


def _apply_paragraph(
    paragraph,
    kind: str,
    sec_idx: int,
    cfg: dict,
    remaining: List[str],
    variant: str = "",
):
    """修正单个页眉/页脚段落格式；修正后复检，残留问题写入 remaining。"""
    expected_alignment = _resolve_alignment_value(cfg.get("alignment"))
    if expected_alignment is not None:
        paragraph.alignment = expected_alignment
    for run in paragraph.runs:
        _apply_run_format(run, cfg)

    leftover: List[str] = []
    _check_paragraph(paragraph, kind, sec_idx, cfg, leftover, variant)
    remaining.extend(leftover)


def _apply_part(
    part, kind: str, sec_idx: int, cfg: dict, remaining: List[str], variant: str = ""
):
    """修正某节的一个页眉/页脚变体；残留问题写入 remaining。"""
    if not _owns_content(part, sec_idx, variant):
        return
    has_part = _has_own_part(part)
    if cfg.get("page_number") and not (has_part and _has_page_field(part)):
        _insert_page_field(part, cfg)
        has_part = True  # 插入页码域时部件已随之创建
        if not _has_page_field(part):
            remaining.append(_page_number_issue(kind, sec_idx, variant))
    paragraphs = _content_paragraphs(part) if has_part else []
    if not paragraphs:
        _check_missing_content(kind, sec_idx, variant, cfg, remaining)
        return
    for p in paragraphs:
        _apply_paragraph(p, kind, sec_idx, cfg, remaining, variant)


def apply_header_footer(document, config_model=None) -> List[str]:
    """按配置修正文档所有节的页眉页脚与页面设置。

    Returns:
        修正后仍残留的问题文案列表（空表示全部修正或未启用）。
    """
    hf_config = resolve_hf_config(config_model)
    if not hf_config:
        return []
    remaining: List[str] = []
    page_cfg = hf_config.get("page_setup") or {}
    even_pages_on = _setting_on(document, "w:evenAndOddHeaders")
    for idx, section in enumerate(document.sections, start=1):
        for kind, cfg, part, variant in _section_parts(
            section, idx, hf_config, even_pages_on
        ):
            _apply_part(part, kind, idx, cfg, remaining, variant)
        if page_cfg:
            _apply_page_setup(section, idx, page_cfg, remaining)
    _apply_doc_flags(document, hf_config, remaining)
    if remaining:
        logger.debug(f"页眉页脚/页面设置修正完成，残留 {len(remaining)} 个问题")
    return remaining
