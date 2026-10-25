"""正文标点规则（rules/body.py）：半角标点检测与 run 拆分。"""

from unittest.mock import patch

import pytest
from docx import Document

from wordformat.rules.body import (
    _HALF_TO_FULL,
    BodyText,
    _split_run_at,
)


def _node(text):
    doc = Document()
    p = doc.add_paragraph(text)
    node = BodyText(value={"category": "body_text"}, level=1, paragraph=p)
    node.load_config({})
    return doc, node


def _issues(node, doc):
    with patch.object(node, "add_comment") as mock_comment:
        node._check_punctuation(doc, {})
    return mock_comment.call_args_list


# 独立于被测实现的期望映射：实现里的常量改错时，这里必须失败
_EXPECTED_PUNCT = {
    ",": "，",
    ".": "。",
    ";": "；",
    ":": "：",
    "(": "（",
    ")": "）",
    "[": "【",
    "]": "】",
    '"': "“",
    "'": "‘",
    "!": "！",
    "?": "？",
}


class TestPunctuationMapping:
    def test_mapping_table_unchanged(self):
        assert _HALF_TO_FULL == _EXPECTED_PUNCT

    @pytest.mark.parametrize(("half", "full"), sorted(_EXPECTED_PUNCT.items()))
    def test_half_width_mapped_to_full_width(self, half, full):
        doc, node = _node(f"中{half}文")
        calls = _issues(node, doc)
        assert len(calls) == 1
        text = calls[0].kwargs["text"]
        assert f'"{half}(英文)"' in text
        assert f'"{full}(中文)"' in text

    def test_start_and_end_of_paragraph(self):
        doc, node = _node(",中文,中文.")
        calls = _issues(node, doc)
        assert len(calls) == 3

    def test_punctuation_in_chinese_context_only(self):
        # 英文句子里的半角标点不应误报（前后不是中文）
        doc, node = _node("hello, world. 3.14")
        assert _issues(node, doc) == []


class TestPunctuationTargets:
    def test_comment_anchored_to_punctuation_run(self):
        doc, node = _node("中文,中文")
        calls = _issues(node, doc)
        assert calls[0].kwargs["runs"].text == ","
        assert node.paragraph.text == "中文,中文"

    def test_punctuation_split_across_runs(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("中文")
        p.add_run(",中文")
        node = BodyText(value={"category": "body_text"}, level=1, paragraph=p)
        node.load_config({})
        calls = _issues(node, doc)
        assert calls[0].kwargs["runs"].text == ","
        assert p.text == "中文,中文"

    def test_citation_internal_punctuation_skipped(self):
        doc, node = _node("文献[1,2]表明,该方法有效。")
        calls = _issues(node, doc)
        assert len(calls) == 1
        assert calls[0].kwargs["runs"].text == ","

    def test_missing_split_falls_back_to_first_run(self):
        doc, node = _node("中文,中文")
        with patch("wordformat.rules.body._split_run_at", return_value=None):
            calls = _issues(node, doc)
        assert calls[0].kwargs["runs"]._element is node.paragraph.runs[0]._element

    def test_no_paragraph_skipped(self):
        node = BodyText(value={"category": "body_text"}, level=1, paragraph=None)
        _issues(node, Document())  # 不抛异常

    def test_rule_wiring_reports_punctuation(self):
        doc, node = _node("中文,中文")
        with patch.object(node, "add_comment") as mock_comment:
            node.check_format(doc)
        texts = [call.kwargs["text"] for call in mock_comment.call_args_list]
        assert any("标点问题" in text for text in texts)


class TestSplitRunAt:
    def test_split_returns_run_with_requested_range(self):
        doc = Document()
        p = doc.add_paragraph("中文,中文")
        run = _split_run_at(p, 2, 3)
        assert run is not None
        assert run.text == ","
        assert p.text == "中文,中文"

    def test_out_of_range_returns_none(self):
        doc = Document()
        p = doc.add_paragraph("中文")
        assert _split_run_at(p, 10, 11) is None


class TestApplyReplaceClearsCitationFormat:
    def test_superscript_removed_after_replace(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("旧[1]").font.superscript = True
        p.add_run("文本")
        node = BodyText(
            value={"category": "body_text", "replace": "新文本"},
            level=1,
            paragraph=p,
        )
        assert node.apply_replace(doc) is True
        assert p.text == "新文本"
        assert all(run.font.superscript is None for run in p.runs)

    def test_no_replace_keeps_superscript(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("旧[1]").font.superscript = True
        node = BodyText(value={"category": "body_text"}, level=1, paragraph=p)
        assert node.apply_replace(doc) is False
        assert p.runs[0].font.superscript is True


def test_punctuation_run_keeps_run_formatting():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("中文")
    p.add_run(",中文").bold = True
    node = BodyText(value={"category": "body_text"}, level=1, paragraph=p)
    node.load_config({})
    calls = _issues(node, doc)
    punctuation = calls[0].kwargs["runs"]
    assert punctuation.text == ","
    assert punctuation.bold is True
