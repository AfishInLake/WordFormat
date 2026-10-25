#! /usr/bin/env python
# @Time    : 2026/10/24
# @Author  : afish
# @File    : test_gbt7714.py
"""GB/T 7714 著录格式校验（rules/references.py）。"""

from unittest.mock import patch

import pytest
from docx import Document
from docx.shared import Pt

from wordformat.domains.thesis.rules.references import (
    _GBT_DOC_TYPES,
    _TYPE_CHECKS,
    ReferenceEntry,
    check_gbt7714_entry,
)

# GB/T 7714 规则可能报出的问题类型；段落/字符样式批注用同一个位置标签，
# 所以集成测试必须按问题类型过滤，而不是按「参考文献条目-」前缀。
_GBT_LABELS = {
    "类型标识缺失",
    "类型标识错误",
    "出版年缺失",
    "卷期页缺失",
    "页码缺失",
    "出版项缺失",
    "出版日期缺失",
    "报纸名缺失",
    "版次缺失",
    "标准号缺失",
    "专利号缺失",
    "公告日期缺失",
    "引用日期缺失",
    "获取路径缺失",
    "著者省略错误",
    "结尾符缺失",
    "标点问题",
}

# 合规样例：类型 -> 合法著录。同一类型可有多条合法变体，但文本不得重复。
_COMPLIANT: dict[str, tuple[str, ...]] = {
    "J": (
        "[1] 张三, 李四, 王五. 某某问题研究[J]. 计算机学报, 2020, 43(5): 100-110.",
        "[2] 赵某. 算法分析[J]. 软件学报, 2021(6): 33-40.",
        "[3] Smith J, Brown A, Taylor B, et al. Deep learning for NLP[J]. "
        "IEEE Transactions on Neural Networks and Learning Systems, "
        "2019, 30(2): 45-56.",
        "[4] 张三, 李四, 王五, 等. 某问题再研究[J]. 计算机学报, 2022, 45(1): 1-9.",
        "[5] 王五. 在线优先发表研究[J/OL]. 计算机学报, 2020, 43(5)"
        "[2021-01-01]. https://example.com.",
    ),
    "M": (
        "[6] 王某. 机器学习[M]. 北京: 清华大学出版社, 2019.",
        "[7] 萧钰. 出版业信息化迈入快车道[M/OL]. (2001-12-19)[2002-04-15]. "
        "http://www.creader.com/news/20011219/200112190019.html.",
    ),
    "C": (
        "[8] 李某. 某算法[C]//张三. 论文集. 北京: 机械工业出版社, 2018: 12-20.",
        "[9] 中国计算机学会. 论文集[C]. 北京: 机械工业出版社, 2018.",
    ),
    "G": ("[10] 国家统计局. 中国统计年鉴[G]. 北京: 中国统计出版社, 2021.",),
    "N": (
        "[11] 张三. 某新闻[N]. 人民日报, 2020-05-01(3).",
        "[12] 李四. 某评论[N]. 光明日报, 2020-05-01(2).",
    ),
    "D": ("[13] 李某. 系统设计[D]. 北京: 北京大学, 2021.",),
    "R": (
        "[14] 中国科技发展战略研究小组. 中国科技发展研究报告[R]. "
        "北京: 科学出版社, 2020.",
    ),
    "S": (
        "[15] 国家质量监督检验检疫总局. "
        "GB/T 7713.3—2014 科技报告编写规则[S]. 北京: 中国标准出版社, 2015.",
    ),
    "P": (
        "[16] 张三. 一种数据处理装置: CN101234567A[P]. 2015-06-01.",
        "[17] 李四. 一种装置: ZL 2015 1 0123456.7[P]. 2015-06-01.",
    ),
    "DB": (
        "[18] 国家统计局. 中国统计数据库[DB/CD]. 北京: 中国统计出版社, 2020.",
        "[19] 机构. 数据库[DB/OL]. (2020-01-01)[2021-01-01]. https://example.com.",
    ),
    "CP": (
        "[20] 某公司. 分析软件[CP/CD]. 北京: 某出版社, 2019.",
        "[21] 机构. 软件[CP/OL]. (2020-01-01)[2021-01-01]. https://example.com.",
    ),
    "EB": (
        "[22] 中国互联网络信息中心. 第49次中国互联网络发展状况统计报告[EB/OL]. "
        "(2022-01-25)[2022-03-10]. https://www.cnnic.net.cn/n4/index.html.",
        "[23] 机构. 电子公告[EB/CD]. 北京: 某出版社, 2020.",
    ),
}

# 违规样例：类型 -> ((条目, 期望问题类型), ...)，每种类型至少一条。
_VIOLATIONS: dict[str, tuple[tuple[str, tuple[str, ...]], ...]] = {
    "J": (
        ("[1] 张三. 研究[J]. 计算机学报, 43(5): 100-110.", ("出版年缺失",)),
        ("[2] 张三. 研究[J]. 计算机学报, 2020.", ("卷期页缺失",)),
        ("[3] 张三. 研究[J]. 计算机学报, 2020, 43(5).", ("页码缺失",)),
        (
            "[4] 张三. 研究[J/OL]. 计算机学报, 43(5): 100-110[2021-01-01]. "
            "https://example.com.",
            ("出版年缺失",),
        ),
    ),
    "M": (("[4] 王某. 机器学习[M]. 2019.", ("出版项缺失",)),),
    "C": (("[5] 李某. 某算法[C]//张三. 论文集.", ("出版项缺失",)),),
    "G": (("[6] 国家统计局. 中国统计年鉴[G]. 2021.", ("出版项缺失",)),),
    "N": (
        ("[7] 张三. 某新闻[N]. 人民日报(3).", ("出版日期缺失",)),
        ("[8] 张三. 某新闻[N]. 人民日报, 2020-05-01.", ("版次缺失",)),
        ("[9] 张三. 某新闻[N]. 2020-05-01(3).", ("报纸名缺失",)),
        ("[10] 张三. 某新闻[N]. 人民日报, 2020年5月1日(3).", ("出版日期缺失",)),
        ("[11] 张三. 某新闻[N]. 人民日报, [2021-01-01](3).", ("出版日期缺失",)),
    ),
    "D": (("[12] 李某. 系统设计[D]. 2021.", ("出版项缺失",)),),
    "R": (("[13] 某课题组. 研究报告[R]. 2020.", ("出版项缺失",)),),
    "S": (
        (
            "[14] 国家质量监督检验检疫总局. 科技报告编写规则[S]. "
            "北京: 中国标准出版社, 2015.",
            ("标准号缺失",),
        ),
        (
            "[15] 国家质量监督检验检疫总局. GB/T 7713.3—2014 科技报告编写规则[S].",
            ("出版项缺失",),
        ),
    ),
    "P": (
        ("[16] 张三. 一种数据处理装置[P]. 2015-06-01.", ("专利号缺失",)),
        ("[17] 张三. 一种装置: CN101234567A[P]. 北京.", ("公告日期缺失",)),
        ("[18] 张三. 一种装置: CN101234567A[P]. 2015年6月1日.", ("公告日期缺失",)),
        (
            "[19] 张三. 一种装置: CN101234567A[P]. [2015-06-01].",
            ("公告日期缺失",),
        ),
    ),
    "DB": (("[20] 国家统计局. 统计数据库[DB/CD]. 2020.", ("出版项缺失",)),),
    "CP": (("[21] 某公司. 分析软件[CP/CD]. 2020.", ("出版项缺失",)),),
    "EB": (("[22] 机构. 电子公告[EB/CD]. 2020.", ("出版项缺失",)),),
}

_COMPLIANT_CASES = [(t, text) for t, texts in _COMPLIANT.items() for text in texts]
_VIOLATION_CASES = [
    (t, text, expected) for t, cases in _VIOLATIONS.items() for text, expected in cases
]
_DISABLE_CASES = (
    ("check_journal_item", "[1] 张三. 研究[J]. 计算机学报."),
    ("check_publisher_item", "[1] 王某. 机器学习[M]. 2019."),
    ("check_newspaper_item", "[1] 张三. 新闻[N]. 人民日报(3)."),
    ("check_standard_item", "[1] 某局. 规则[S]."),
    ("check_patent_item", "[1] 张三. 装置[P]. 北京."),
)


def _case_ids(cases):
    """用例 id 形如 J-1、N-3，避免 pytest 输出整条著录。"""
    counters: dict[str, int] = {}
    ids = []
    for case in cases:
        doc_type = case[0]
        counters[doc_type] = counters.get(doc_type, 0) + 1
        ids.append(f"{doc_type}-{counters[doc_type]}")
    return ids


def _labels(text, cfg=None):
    """只取问题类型列表，便于断言。"""
    return [issue[0] for issue in check_gbt7714_entry(text, cfg or {})]


def _gbt_issues(comment_texts):
    """从批注文本列表中提取 GB/T 7714 问题类型，忽略样式类批注。"""
    found = []
    for text in comment_texts:
        label = text.split("：", 1)[0].split("-")[-1]
        if label in _GBT_LABELS:
            found.append(label)
    return found


class TestTypeCoverage:
    """12 种文献类型必须都有校验实现，且都有正反用例。"""

    def test_every_gbt_type_has_field_checks(self):
        assert set(_TYPE_CHECKS) == _GBT_DOC_TYPES
        assert all(_TYPE_CHECKS[t] for t in _GBT_DOC_TYPES)

    def test_every_gbt_type_has_compliant_case(self):
        assert set(_COMPLIANT) == _GBT_DOC_TYPES

    def test_every_gbt_type_has_violation_case(self):
        assert set(_VIOLATIONS) == _GBT_DOC_TYPES

    def test_cases_are_unique(self):
        texts = [text for _, text in _COMPLIANT_CASES]
        texts += [text for _, text, _ in _VIOLATION_CASES]
        assert len(texts) == len(set(texts))


class TestCompliantEntries:
    """规范条目必须零报告，否则规则会淹没真实问题。"""

    @pytest.mark.parametrize(
        "doc_type,text", _COMPLIANT_CASES, ids=_case_ids(_COMPLIANT_CASES)
    )
    def test_compliant_entry_has_no_issue(self, doc_type, text):
        assert _labels(text) == []


class TestViolations:
    @pytest.mark.parametrize(
        "doc_type,text,expected",
        _VIOLATION_CASES,
        ids=_case_ids(_VIOLATION_CASES),
    )
    def test_reports_expected_issue(self, doc_type, text, expected):
        assert _labels(text) == list(expected)

    @pytest.mark.parametrize("code", ["DB/OL", "CP/OL", "EB/OL"])
    def test_online_carrier_requires_date_and_path(self, code):
        text = f"[1] 机构. 资料[{code}]. 北京: 某出版社, 2020."
        assert _labels(text) == ["引用日期缺失", "获取路径缺失"]

    @pytest.mark.parametrize("cfg_key,text", _DISABLE_CASES)
    def test_type_check_can_be_disabled(self, cfg_key, text):
        assert _labels(text, {cfg_key: False}) == []


class TestSampleRegressions:
    """审查样例时发现的误报/漏报：样例合法或日期写法不同，规则必须给出正确结论。"""

    def test_online_monograph_without_publisher_is_compliant(self):
        # GB/T 7714 经典 [M/OL] 样例：联机专著可省略出版项
        text = (
            "[1] 萧钰. 出版业信息化迈入快车道[M/OL]. (2001-12-19)[2002-04-15]. "
            "http://www.creader.com/news/20011219/200112190019.html."
        )
        assert _labels(text) == []

    def test_online_monograph_still_requires_citation_date_and_path(self):
        text = "[1] 萧钰. 题名[M/OL]. 北京: 某出版社, 2020."
        assert _labels(text) == ["引用日期缺失", "获取路径缺失"]

    def test_publication_year_not_satisfied_by_citation_date(self):
        # [引用日期] 里的年份不能顶替出版年
        text = (
            "[1] 张三. 研究[J/OL]. 计算机学报, 43(5): 100-110[2021-01-01]. "
            "https://example.com."
        )
        assert _labels(text) == ["出版年缺失"]

    def test_online_journal_may_omit_pages(self):
        text = (
            "[1] 王五. 在线研究[J/OL]. 计算机学报, 2020, 43(5)"
            "[2021-01-01]. https://example.com."
        )
        assert _labels(text) == []

    def test_volume_only_journal_is_compliant(self):
        # 只有卷号、没有期号是合法写法
        text = "[1] Smith J. Study[J]. Journal of Testing, 2019, 30: 45-56."
        assert _labels(text) == []

    def test_newspaper_requires_iso_date(self):
        text = "[1] 张三. 新闻[N]. 人民日报, 2020年5月1日(3)."
        assert _labels(text) == ["出版日期缺失"]

    def test_patent_requires_iso_date(self):
        text = "[1] 张三. 装置: CN101234567A[P]. 2015年6月1日."
        assert _labels(text) == ["公告日期缺失"]


class TestTypeMarker:
    def test_missing_marker(self):
        assert _labels("[1] 张三. 某某研究. 计算机学报, 2020.") == ["类型标识缺失"]

    def test_illegal_marker(self):
        assert _labels("[1] 张三. 某某研究[K]. 计算机学报, 2020.") == ["类型标识错误"]

    def test_numeric_serial_not_treated_as_marker(self):
        # [1] 是条目序号，不能被当成类型标识
        assert "类型标识错误" not in _labels("[1] 张三. 研究. 某刊, 2020.")

    def test_marker_check_can_be_disabled(self):
        assert (
            _labels(
                "[1] 张三. 某某研究. 计算机学报, 2020.", {"require_type_marker": False}
            )
            == []
        )


class TestAuthorsAndOptionalChecks:
    def test_ellipsis_reported(self):
        assert _labels("[1] 张三, 李四, …… . 研究[J]. 某刊, 2020, 1(2): 3.") == [
            "著者省略错误"
        ]

    def test_chinese_entry_with_et_al(self):
        text = "[1] 张三, 李四, et al. 研究[J]. 某刊, 2020, 1(2): 3."
        assert _labels(text) == ["著者省略错误"]

    def test_ellipsis_can_be_disabled(self):
        text = "[1] 张三, …… . 研究[J]. 某刊, 2020, 1(2): 3."
        assert _labels(text, {"check_author_ellipsis": False}) == []

    def test_terminator_off_by_default(self):
        text = "[1] 王某. 机器学习[M]. 北京: 清华大学出版社, 2019"
        assert "结尾符缺失" not in _labels(text)

    def test_terminator_when_enabled(self):
        text = "[1] 王某. 机器学习[M]. 北京: 清华大学出版社, 2019"
        assert _labels(text, {"check_terminator": True}) == ["结尾符缺失"]

    def test_terminator_accepts_trailing_dot(self):
        text = "[1] 王某. 机器学习[M]. 北京: 清华大学出版社, 2019."
        assert _labels(text, {"check_terminator": True}) == []

    def test_punctuation_width_off_by_default(self):
        # 全角括号只触发标点校验，著录项本身仍完整
        text = "[1] 张三. 研究[J]. 计算机学报, 2020, 43（5）: 100-110."
        assert _labels(text) == []

    def test_punctuation_width_when_enabled(self):
        text = "[1] 张三. 研究[J]. 计算机学报, 2020, 43（5）: 100-110."
        assert _labels(text, {"check_punctuation_width": True}) == ["标点问题"]

    def test_punctuation_width_accepts_half_width(self):
        text = "[1] 张三. 研究[J]. 计算机学报, 2020, 43(5): 100-110."
        assert _labels(text, {"check_punctuation_width": True}) == []


class TestNodeIntegration:
    """规则通过 _run_rules 挂载到条目节点，check/apply 都报批注且不改文字。"""

    def _node(self, text):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run(text).font.size = Pt(10.5)
        node = ReferenceEntry(value=p, level=0, paragraph=p)
        node.load_config({})
        return doc, node

    @staticmethod
    def _issues(node, doc, mode):
        with patch.object(node, "add_comment") as mock_comment:
            if mode == "check":
                node.check_format(doc)
            else:
                node.apply_format(doc)
        return [call.kwargs["text"] for call in mock_comment.call_args_list]

    def test_check_reports_gbt_issue(self):
        doc, node = self._node("[1] 张三. 某某研究. 计算机学报, 2020.")
        texts = self._issues(node, doc, "check")
        assert _gbt_issues(texts) == ["类型标识缺失"]

    def test_apply_also_reports_gbt_issue(self):
        doc, node = self._node("[1] 张三. 某某研究. 计算机学报, 2020.")
        texts = self._issues(node, doc, "apply")
        assert _gbt_issues(texts) == ["类型标识缺失"]

    def test_compliant_entry_gets_no_gbt_issue(self):
        doc, node = self._node("[1] 王某. 机器学习[M]. 北京: 清华大学出版社, 2019.")
        texts = self._issues(node, doc, "check")
        assert _gbt_issues(texts) == []

    @pytest.mark.parametrize(
        "text",
        [
            "[1] 张三. 新闻[N]. 人民日报, 2020-05-01(3).",
            "[1] 某局. GB/T 7713.3—2014 规则[S]. 北京: 中国标准出版社, 2015.",
            "[1] 张三. 装置: CN101234567A[P]. 2015-06-01.",
        ],
        ids=["N", "S", "P"],
    )
    def test_other_types_compliant_via_node(self, text):
        doc, node = self._node(text)
        assert _gbt_issues(self._issues(node, doc, "check")) == []

    def test_rule_can_be_disabled_in_config(self):
        doc, node = self._node("[1] 张三. 某某研究. 计算机学报, 2020.")
        node.load_config(
            {"references": {"entry": {"rules": {"gbt_7714": {"enabled": False}}}}}
        )
        assert _gbt_issues(self._issues(node, doc, "check")) == []

    def test_text_not_modified(self):
        text = "[1] 张三. 某某研究. 计算机学报, 2020."
        doc, node = self._node(text)
        node.apply_format(doc)
        assert node.paragraph.text == text

    def test_message_follows_comment_format(self):
        doc, node = self._node("[1] 张三. 某某研究[J]. 计算机学报, 2020.")
        texts = self._issues(node, doc, "check")
        issue = next(t for t in texts if "卷期页缺失" in t)
        assert issue.startswith("参考文献条目-卷期页缺失：当前")
        assert "规范：" in issue

    def test_empty_paragraph_skipped(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("")
        node = ReferenceEntry(value=p, level=0, paragraph=p)
        node.load_config({})
        with patch.object(node, "add_comment") as mock_comment:
            node.check_format(doc)
        texts = [call.kwargs["text"] for call in mock_comment.call_args_list]
        assert _gbt_issues(texts) == []

    def test_paragraph_without_runs_skipped(self):
        doc = Document()
        p = doc.add_paragraph()
        node = ReferenceEntry(value=p, level=0, paragraph=p)
        node.load_config({})
        with patch.object(node, "add_comment") as mock_comment:
            node.check_format(doc)
        mock_comment.assert_not_called()
