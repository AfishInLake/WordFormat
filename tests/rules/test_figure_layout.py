"""图片段落版式规则：环绕方式与超宽检查（rules/object.py）。"""

import base64
from unittest.mock import patch

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

from wordformat.rules.object import FigureImage

# 1x1 透明 PNG，仅用于生成真实的 wp:inline
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _doc_with_picture(tmp_path, width_cm=10.0, wrap=None, behind=False):
    doc = Document()
    p = doc.add_paragraph()
    img = tmp_path / "dot.png"
    img.write_bytes(_PNG)
    p.add_run().add_picture(str(img), width=Cm(width_cm))
    if wrap:
        _set_wrap(p, wrap, behind)
    return doc, p


def _set_wrap(paragraph, wrap_tag: str, behind: bool = False) -> None:
    """把 wp:inline 改写成 wp:anchor，并附上指定环绕元素。"""
    drawing = paragraph._p.xpath(".//w:drawing")[0]
    inline = drawing.find(qn("wp:inline"))
    anchor = OxmlElement("wp:anchor")
    for attr, value in (
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
        ("simplePos", "0"),
        ("relativeHeight", "251658240"),
        ("behindDoc", "1" if behind else "0"),
        ("locked", "0"),
        ("layoutInCell", "1"),
        ("allowOverlap", "1"),
    ):
        anchor.set(attr, value)
    for child in list(inline):
        anchor.append(child)
    if wrap_tag == "wp:wrapNone":
        anchor.append(OxmlElement("wp:wrapNone"))
    else:
        wrap_el = OxmlElement(wrap_tag)
        wrap_el.set("wrapText", "bothSides")
        anchor.append(wrap_el)
    drawing.replace(inline, anchor)


def _node(doc, paragraph):
    node = FigureImage(value={"category": "figure_image"}, level=1, paragraph=paragraph)
    node.load_config({})
    return node


def _issues(node, doc, cfg):
    with patch.object(node, "add_comment") as mock_comment:
        node._check_figure_layout(doc, cfg)
    return [call.kwargs["text"] for call in mock_comment.call_args_list]


class TestWrapCheck:
    def test_inline_matches_embedded(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path)
        assert _issues(_node(doc, p), doc, {"wrap": "嵌入型"}) == []

    def test_inline_reported_when_float_expected(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path)
        texts = _issues(_node(doc, p), doc, {"wrap": "四周型"})
        assert len(texts) == 1
        assert "图片段落-环绕方式错误：当前嵌入型，规范：四周型" in texts[0]

    @pytest.mark.parametrize(
        ("wrap_tag", "label"),
        [
            ("wp:wrapSquare", "四周型"),
            ("wp:wrapTight", "紧密型"),
            ("wp:wrapThrough", "穿越型"),
            ("wp:wrapTopAndBottom", "上下型"),
        ],
    )
    def test_anchor_wrap_variants(self, tmp_path, wrap_tag, label):
        doc, p = _doc_with_picture(tmp_path, wrap=wrap_tag)
        assert _issues(_node(doc, p), doc, {"wrap": label}) == []

    def test_behind_text(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, wrap="wp:wrapNone", behind=True)
        assert _issues(_node(doc, p), doc, {"wrap": "衬于文字下方"}) == []

    def test_behind_text_mismatch(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, wrap="wp:wrapNone", behind=False)
        texts = _issues(_node(doc, p), doc, {"wrap": "衬于文字下方"})
        assert "当前浮于文字上方" in texts[0]

    def test_wrap_accepts_ooxml_value(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, wrap="wp:wrapSquare")
        assert _issues(_node(doc, p), doc, {"wrap": "Square"}) == []

    def test_unknown_wrap_raises(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path)
        with pytest.raises(ValueError, match="环绕方式非法"):
            _issues(_node(doc, p), doc, {"wrap": "环绕"})

    def test_unconfigured_wrap_not_checked(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path)
        assert _issues(_node(doc, p), doc, {}) == []


class TestWidthCheck:
    def test_over_body_width(self, tmp_path):
        # 默认 Letter + 1.25 英寸左右边距，版心 15.24cm
        doc, p = _doc_with_picture(tmp_path, width_cm=20)
        texts = _issues(_node(doc, p), doc, {})
        assert len(texts) == 1
        assert "图片段落-图片超宽" in texts[0]
        assert "20.00cm" in texts[0]
        assert "15.24cm" in texts[0]

    def test_within_body_width(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=10)
        assert _issues(_node(doc, p), doc, {}) == []

    def test_exact_body_width_not_reported(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=15.24)
        assert _issues(_node(doc, p), doc, {}) == []

    def test_explicit_max_width(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=15)
        texts = _issues(_node(doc, p), doc, {"max_width": "12cm"})
        assert "不超过 12.00cm" in texts[0]

    def test_numeric_max_width_is_cm(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=15)
        assert _issues(_node(doc, p), doc, {"max_width": 16}) == []

    def test_missing_unit_raises(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path)
        with pytest.raises(ValueError, match="需带单位"):
            _issues(_node(doc, p), doc, {"max_width": "15"})

    def test_unsupported_unit_raises(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path)
        with pytest.raises(ValueError, match="单位不支持"):
            _issues(_node(doc, p), doc, {"max_width": "15字符"})

    def test_follows_section_margins(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=18)
        section = doc.sections[0]
        section.page_width = Cm(21)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        texts = _issues(_node(doc, p), doc, {})
        assert "不超过 17.00cm" in texts[0]

    def test_multiple_drawings_each_reported(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=20)
        img = tmp_path / "dot.png"
        p.add_run().add_picture(str(img), width=Cm(20))
        texts = _issues(_node(doc, p), doc, {})
        assert len(texts) == 2


class TestRuleWiring:
    def test_rule_enabled_by_default(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=20)
        node = _node(doc, p)
        cfg = node.pydantic_config.rules.figure_layout
        assert cfg.enabled is True
        assert cfg.max_width == "版心"

    def test_disabled_rule_skipped(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=20)
        node = FigureImage(value={"category": "figure_image"}, level=1, paragraph=p)
        node.load_config(
            {"figures": {"image": {"rules": {"figure_layout": {"enabled": False}}}}}
        )
        with patch.object(node, "add_comment") as mock_comment:
            node.check_format(doc)
        assert not any(
            "图片超宽" in call.kwargs["text"] for call in mock_comment.call_args_list
        )

    def test_check_format_reports_width(self, tmp_path):
        doc, p = _doc_with_picture(tmp_path, width_cm=20)
        node = FigureImage(value={"category": "figure_image"}, level=1, paragraph=p)
        node.load_config({"figures": {"image": {"alignment": "居中对齐"}}})
        with patch.object(node, "add_comment") as mock_comment:
            node.check_format(doc)
        texts = [call.kwargs["text"] for call in mock_comment.call_args_list]
        assert any("图片超宽" in t for t in texts)

    def test_no_drawing_no_issue(self):
        doc = Document()
        p = doc.add_paragraph("纯文字段落")
        assert _issues(_node(doc, p), doc, {}) == []
