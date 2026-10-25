#! /usr/bin/env python
# @Time    : 2026/10/22
# @Author  : afish
# @File    : test_gj_tables.py
"""gj 表格导出：table_object 条目（exists / table_rows / 文档序）。"""

from unittest.mock import patch

from docx import Document

from wordformat.base import DocxBase


def _make_docx(tmp_path, body):
    """body: list，元素为 str（段落）或 ("table", rows)。"""
    doc = Document()
    for item in body:
        if isinstance(item, tuple):
            _, rows = item
            num_cols = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=num_cols)
            for i, row in enumerate(rows):
                for j, text in enumerate(row):
                    table.cell(i, j).text = text
        else:
            doc.add_paragraph(item)
    path = tmp_path / "with_tables.docx"
    doc.save(path)
    return str(path)


def _parse(path, n_texts):
    mock_results = [{"label": "body_text", "score": 0.9} for _ in range(n_texts)]
    with patch(
        "wordformat.base.onnx_batch_infer",
        return_value=(mock_results, "body_text"),
    ):
        return DocxBase(path, "/fake/config.yaml").parse()


class TestGjTableExport:
    def test_tables_exported_with_exists(self, tmp_path):
        path = _make_docx(
            tmp_path,
            [
                "标题",
                ("table", [["名称", "数值"], ["A", "1"]]),
                "正文一",
                ("table", [["x", "y"]]),
                "正文二",
            ],
        )
        result = _parse(path, n_texts=3)
        tables = [r for r in result if r["category"] == "table_object"]
        assert len(tables) == 2
        assert tables[0]["table_rows"] == [["名称", "数值"], ["A", "1"]]
        assert tables[1]["table_rows"] == [["x", "y"]]
        assert all(t["exists"] is True for t in tables)
        assert all(t["paragraph"] == "" for t in tables)
        # 文档序：段0, 表1, 段1, 表2, 段2
        assert [r["category"] for r in result] == [
            "body_text",
            "table_object",
            "body_text",
            "table_object",
            "body_text",
        ]
        assert [
            r.get("paragraph") for r in result if r["category"] != "table_object"
        ] == [
            "标题",
            "正文一",
            "正文二",
        ]

    def test_multiple_tables_before_first_paragraph(self, tmp_path):
        path = _make_docx(
            tmp_path,
            [
                ("table", [["a"]]),
                ("table", [["b"]]),
                "正文",
            ],
        )
        result = _parse(path, n_texts=1)
        assert [r["category"] for r in result] == [
            "table_object",
            "table_object",
            "body_text",
        ]
        assert [r["table_rows"][0][0] for r in result[:2]] == ["a", "b"]

    def test_no_tables_keeps_one_to_one(self, tmp_path):
        path = _make_docx(tmp_path, ["A", "B"])
        result = _parse(path, n_texts=2)
        assert len(result) == 2
        assert all(r["category"] != "table_object" for r in result)
