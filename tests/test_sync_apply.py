#! /usr/bin/env python
# @Time    : 2026/10/22
# @Author  : afish
# @File    : test_sync_apply.py
"""af --sync 结构同步单元测试（sync_apply.py）。"""

import base64
import json

import pytest
from docx import Document

from wordformat.pipeline.orchestrate import auto_format_thesis_document
from wordformat.pipeline.sync_apply import (
    IMAGE_SENTINEL,
    _item_key,
    _item_text,
    sync_document_structure,
)

_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _item(category, text, **extra):
    d = {"category": category, "paragraph": text, "score": 1.0}
    d.update(extra)
    return d


def _make_doc(texts):
    doc = Document()
    for t in texts:
        doc.add_paragraph(t)
    return doc


class TestHelpers:
    def test_item_text_prefers_replace(self):
        assert _item_text({"paragraph": "旧", "replace": "新"}) == "新"

    def test_item_text_figure_image_returns_empty(self):
        # figure_image 的 paragraph 是路径，不可写入文档
        assert _item_text({"category": "figure_image", "paragraph": "/a.png"}) == ""

    def test_item_text_table_returns_empty(self):
        # table_object 不含可写入段落的文本
        assert _item_text({"category": "table_object", "paragraph": "x"}) == ""

    def test_item_text_empty_replace_is_authoritative(self):
        # replace: "" 表示清空，不能被 or 回退成 paragraph
        assert _item_text({"paragraph": "旧文本", "replace": ""}) == ""

    def test_item_key_sentinel_for_figure(self):
        assert _item_key(_item("figure_image", "/a.png")) == IMAGE_SENTINEL


class TestSyncDocumentStructure:
    def test_delete_removes_paragraph(self):
        doc = _make_doc(["A", "B", "C"])
        sync_document_structure(doc, [_item("body_text", "A"), _item("body_text", "C")])
        assert [p.text for p in doc.paragraphs] == ["A", "C"]

    def test_insert_middle(self):
        doc = _make_doc(["A", "C"])
        sync_document_structure(
            doc,
            [_item("body_text", "A"), _item("body_text", "B"), _item("body_text", "C")],
        )
        assert [p.text for p in doc.paragraphs] == ["A", "B", "C"]

    def test_insert_multiple_keeps_order(self):
        doc = _make_doc(["A", "Z"])
        items = (
            [_item("body_text", "A")]
            + [_item("body_text", f"N{i}") for i in range(3)]
            + [_item("body_text", "Z")]
        )
        sync_document_structure(doc, items)
        assert [p.text for p in doc.paragraphs] == ["A", "N0", "N1", "N2", "Z"]

    def test_insert_at_tail(self):
        doc = _make_doc(["A"])
        sync_document_structure(
            doc, [_item("body_text", "A"), _item("body_text", "尾部新段")]
        )
        assert [p.text for p in doc.paragraphs] == ["A", "尾部新段"]

    def test_text_edit_applied(self):
        doc = _make_doc(["原文", "保持"])
        sync_document_structure(
            doc, [_item("body_text", "改后"), _item("body_text", "保持")]
        )
        assert [p.text for p in doc.paragraphs] == ["改后", "保持"]

    def test_replace_field_edits_text(self):
        doc = _make_doc(["原文"])
        sync_document_structure(doc, [_item("body_text", "原文", replace="替换文")])
        assert doc.paragraphs[0].text == "替换文"

    def test_empty_replace_clears_text(self):
        doc = _make_doc(["原文"])
        sync_document_structure(doc, [_item("body_text", "原文", replace="")])
        assert doc.paragraphs[0].text == ""

    def test_image_sentinel_survives_neighbors_edit(self, tmp_path):
        # 文档中间有真实图片段（空文本），在其前后增删，图片段不得被删/覆写
        doc = _make_doc(["前段"])
        img_para = doc.add_paragraph()
        img_para.add_run().add_picture(str(_write_png(tmp_path)))
        doc.add_paragraph("后段")
        items = [
            _item("body_text", "前段改"),
            _item("figure_image", ""),
            _item("body_text", "新增段"),
            _item("body_text", "后段"),
        ]
        sync_document_structure(doc, items)
        from wordformat.utils import para_contains_image

        images = [p for p in doc.paragraphs if para_contains_image(p)]
        assert len(images) == 1
        assert [p.text for p in doc.paragraphs if not para_contains_image(p)] == [
            "前段改",
            "新增段",
            "后段",
        ]

    def test_replace_with_more_items_keeps_json_order(self):
        # B=['A','X'] ← A=['A','B','C']：多出的条目须排在配对覆写段之后
        doc = _make_doc(["A", "X"])
        sync_document_structure(
            doc,
            [_item("body_text", "A"), _item("body_text", "B"), _item("body_text", "C")],
        )
        assert [p.text for p in doc.paragraphs] == ["A", "B", "C"]

    def test_replace_with_more_items_before_following_equal(self):
        # replace 区段后还有 equal 段时，多出的条目不能插到区首导致顺序倒置
        doc = _make_doc(["A", "X", "Z"])
        sync_document_structure(
            doc,
            [
                _item("body_text", "A"),
                _item("body_text", "B"),
                _item("body_text", "C"),
                _item("body_text", "Z"),
            ],
        )
        assert [p.text for p in doc.paragraphs] == ["A", "B", "C", "Z"]

    def test_delete_all_keeps_one_paragraph(self):
        doc = _make_doc(["只有一段"])
        sync_document_structure(doc, [])
        # Word 要求正文至少一个段落
        assert len(doc.paragraphs) == 1
        assert doc.paragraphs[0].text == ""


def _write_png(tmp_path):
    p = tmp_path / "fig.png"
    p.write_bytes(_PNG_1PX)
    return p


class TestAfSyncIntegration:
    def _run(self, tmp_path, texts, items, sync):
        dp = tmp_path / "s.docx"
        _make_doc(texts).save(dp)
        jp = tmp_path / "s.json"
        jp.write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
        return auto_format_thesis_document(
            jsonpath=str(jp),
            docxpath=str(dp),
            check=False,
            savepath=str(tmp_path / "out"),
            sync=sync,
        )

    def test_sync_applies_add_delete_edit(self, tmp_path):
        out = self._run(
            tmp_path,
            ["绪论", "正文A", "正文B"],
            [
                _item("heading_level_1", "绪论"),
                _item("body_text", "正文A改"),
                _item("body_text", "新插入段"),
            ],
            sync=True,
        )
        doc = Document(out)
        assert [p.text for p in doc.paragraphs] == ["绪论", "正文A改", "新插入段"]

    def test_sync_replace_with_extra_items_keeps_order(self, tmp_path):
        # 编辑 JSON 在替换段中间新增一段：输出顺序须与 JSON 完全一致
        out = self._run(
            tmp_path,
            ["绪论", "正文A", "正文B"],
            [
                _item("heading_level_1", "绪论"),
                _item("body_text", "正文A改"),
                _item("body_text", "新增段"),
                _item("body_text", "正文B"),
            ],
            sync=True,
        )
        doc = Document(out)
        assert [p.text for p in doc.paragraphs] == [
            "绪论",
            "正文A改",
            "新增段",
            "正文B",
        ]

    def test_sync_empty_json_keeps_placeholder(self, tmp_path):
        # 编辑 JSON 删空全部内容：Word 正文保留一个空段，流程不得因计数不一致报错
        out = self._run(tmp_path, ["A", "B"], [], sync=True)
        doc = Document(out)
        assert [p.text for p in doc.paragraphs] == [""]

    def test_no_sync_count_mismatch_raises(self, tmp_path):
        # 段落数与条目数不等且未开 sync：维持现状的显式报错保护
        with pytest.raises(ValueError, match="不一致"):
            self._run(
                tmp_path,
                ["A", "B", "C"],
                [_item("body_text", "A"), _item("body_text", "C")],
                sync=False,
            )

    def test_no_sync_empty_replace_clears_text(self, tmp_path):
        # 非 sync 路径的 apply_replace 也要把 replace:"" 视为清空
        out = self._run(
            tmp_path,
            ["绪论", "原文本"],
            [
                _item("heading_level_1", "绪论"),
                _item("body_text", "原文本", replace=""),
            ],
            sync=False,
        )
        assert Document(out).paragraphs[1].text == ""

    def test_table_item_does_not_touch_paragraphs(self, tmp_path):
        # JSON 中的原表格条目（exists）不参与对齐，段落文本不受影响
        out = self._run(
            tmp_path,
            ["A", "B"],
            [
                _item("body_text", "A"),
                _item("table_object", "", table_rows=[["x"]], exists=True),
                _item("body_text", "B"),
            ],
            sync=True,
        )
        doc = Document(out)
        assert [p.text for p in doc.paragraphs] == ["A", "B"]

    def test_new_table_item_ignored_in_sync(self, tmp_path):
        # sync 不支持新增表格：条目标 sync_skip 并被 zip/建表跳过，段落不变
        out = self._run(
            tmp_path,
            ["A"],
            [
                _item("body_text", "A"),
                _item("table_object", "", table_rows=[["x"]]),
            ],
            sync=True,
        )
        doc = Document(out)
        assert [p.text for p in doc.paragraphs] == ["A"]
        assert doc.tables == []
