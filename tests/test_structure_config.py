#! /usr/bin/env python
# @Time    : 2026/10/22
# @Author  : afish
# @File    : test_structure_config.py
"""结构类别策略配置化：terminal / void / skip_children 按名单驱动。"""

from docx import Document

from wordformat.pipeline.stages import FormattingExecutionStage
from wordformat.rules.node import FormatNode
from wordformat.settings import (
    DEFAULT_SKIP_CHILDREN_CATEGORIES,
    VOIDNODELIST,
    resolve_structure_categories,
)


class _RecNode(FormatNode):
    """记录是否被格式化的节点。"""

    calls: list = []

    def check_format(self, doc):
        self.calls.append(self.value.get("category"))


def _build_tree():
    root = FormatNode(value={"category": "top", "paragraph": ""}, level=0)
    heading = FormatNode(
        value={"category": "heading_level_1", "paragraph": "第一章"}, level=1
    )
    body = _RecNode(value={"category": "body_text", "paragraph": "正文"}, level=999)
    fulu = _RecNode(value={"category": "heading_fulu", "paragraph": "附录"}, level=1)
    child = _RecNode(
        value={"category": "body_text", "paragraph": "附录内容"}, level=999
    )
    root.add_child_node(heading)
    heading.add_child_node(body)
    root.add_child_node(fulu)
    fulu.add_child_node(child)

    for node in (heading, body, fulu, child):
        node.paragraph = Document().add_paragraph("x")
    return root, child


def _run(root, void=None, skip=None):
    _RecNode.calls = []
    doc = Document()
    FormattingExecutionStage().apply_format_check_to_all_nodes(
        root,
        doc,
        {},
        check=True,
        void_categories=void,
        skip_children_categories=skip,
    )
    return list(_RecNode.calls)


class TestResolveStructureCategories:
    def test_defaults_when_no_config(self):
        resolved = resolve_structure_categories({})
        assert resolved["void"] == VOIDNODELIST
        assert resolved["skip_children"] == DEFAULT_SKIP_CHILDREN_CATEGORIES

    def test_yaml_overrides(self):
        cfg = {
            "structure": {
                "void_categories": ["top"],
                "skip_children_categories": ["other"],
            }
        }
        resolved = resolve_structure_categories(cfg)
        assert resolved["void"] == ["top"]
        assert resolved["skip_children"] == ["other"]

    def test_empty_list_kept_as_empty(self):
        """显式空列表是权威的：不排除任何类别，不回退到内置默认值。"""
        cfg = {
            "structure": {
                "void_categories": [],
                "skip_children_categories": [],
            }
        }
        resolved = resolve_structure_categories(cfg)
        assert resolved["void"] == []
        assert resolved["skip_children"] == []

    def test_invalid_value_falls_back_to_default(self):
        cfg = {"structure": {"void_categories": "not-a-list"}}
        resolved = resolve_structure_categories(cfg)
        assert resolved["void"] == VOIDNODELIST


class TestTraverseConfig:
    def test_default_appendix_subtree_skipped(self):
        root, _ = _build_tree()
        calls = _run(root)
        # 附录标题命中 void，子树 skip_children，附录子节点不参与格式化
        assert calls.count("body_text") == 1
        assert "heading_fulu" not in calls

    def test_appendix_included_when_removed_from_lists(self):
        root, _ = _build_tree()
        void = [c for c in VOIDNODELIST if c not in ("heading_fulu", "other")]
        skip = [c for c in DEFAULT_SKIP_CHILDREN_CATEGORIES if c != "heading_fulu"]
        calls = _run(root, void=void, skip=skip)
        # 附录标题参与格式化，其子树也被遍历（两次 body_text）
        assert "heading_fulu" in calls
        assert calls.count("body_text") == 2

    def test_other_category_excluded_by_void_list(self):
        root = FormatNode(value={"category": "top"}, level=0)
        other = _RecNode(value={"category": "other", "paragraph": "封面"}, level=999)
        other.paragraph = Document().add_paragraph("封面")
        root.add_child_node(other)
        calls = _run(root)
        assert "other" not in calls
        calls = _run(root, void=["top"])
        assert "other" in calls

    def test_empty_lists_disable_all_exclusions(self):
        """void/skip_children 显式空列表：所有节点都参与格式化且子树全遍历。"""
        root, _ = _build_tree()
        calls = _run(root, void=[], skip=[])
        assert "heading_fulu" in calls
        assert calls.count("body_text") == 2
