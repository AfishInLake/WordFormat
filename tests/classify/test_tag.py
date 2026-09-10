"""
Core 模块综合测试

覆盖 tree.py, utils.py, rules/node.py, numbering.py, settings.py
"""

import os
import pytest
from io import StringIO
from unittest.mock import MagicMock, patch

from docx import Document
from docx.oxml.ns import qn

from wordformat.tree import Tree, Stack, print_tree
from wordformat.rules.node import TreeNode, FormatNode
from wordformat.numbering import (
    _auto_strip_numbering,
    _strip_reference_numbering,
    apply_auto_numbering,
    create_numbering_definition,
    process_heading_numbering,
)
from wordformat.utils import (
    get_file_name,
    ensure_is_directory,
    ensure_directory_exists,
    _to_roman,
    _to_chinese_num,
    load_yaml_with_merge,
    get_paragraph_numbering_text,
    remove_all_numbering,
    _format_number,
    _get_level_fmt,
    _count_numbering_levels,
)
from wordformat.base import (
    DocxBase,
    _apply_footer,
    _apply_section_state,
    _apply_threshold,
    _fix_abstract_en_title,
    _fix_abstract_titles,
    _fix_document_title,
    _fix_known_categories,
    _fix_sequence,
    _fix_toc,
    _neutralize_appendix,
)
from wordformat import settings


# ============================================================
# tree.py — Tree
# ============================================================


# ============================================================
# base.py — DocxBase
# ============================================================


class TestDocxBase:
    """测试 DocxBase 类的初始化和 parse 方法"""

    def test_init(self, temp_docx):
        """测试 DocxBase 初始化"""
        base = DocxBase(temp_docx, "/fake/config.yaml")
        assert base.docx_file == temp_docx
        assert base.document is not None
        assert base.re_dict == {}

    def _create_multi_para_docx(self, tmp_path, texts):
        """辅助方法：创建包含多个段落的 docx 文件"""
        doc = Document()
        for text in texts:
            doc.add_paragraph(text)
        path = str(tmp_path / "multi.docx")
        doc.save(path)
        return path

    def test_parse_with_mocked_batch_infer(self, tmp_path):
        """测试 parse 方法使用 mock 的批量推理"""
        path = self._create_multi_para_docx(tmp_path, ["绪论", "研究背景", "正文内容"])
        mock_batch_results = [
            {"label": "heading_level_1", "score": 0.95},
            {"label": "body_text", "score": 0.88},
            {"label": "body_text", "score": 0.75},
        ]

        with patch("wordformat.base.onnx_batch_infer", return_value=mock_batch_results):
            base = DocxBase(path, "/fake/config.yaml")
            result = base.parse()

        assert len(result) == 3
        assert result[0]["category"] == "heading_level_1"
        assert result[1]["category"] == "body_text"
        assert "paragraph" in result[0]
        assert "score" in result[0]

    def test_parse_low_score_keeps_label_and_flags_review(self, tmp_path):
        """C2：低置信度不再硬砍为 body_text，而是保留原标签并标记 needs_review"""
        path = self._create_multi_para_docx(tmp_path, ["绪论", "研究背景"])
        mock_batch_results = [
            {"label": "heading_level_1", "score": 0.95},
            {"label": "heading_level_2", "score": 0.3},
        ]

        with patch("wordformat.base.onnx_batch_infer", return_value=mock_batch_results):
            base = DocxBase(path, "/fake/config.yaml")
            result = base.parse()

        assert result[0]["category"] == "heading_level_1"
        assert result[0]["needs_review"] is False
        # 低置信度：保留 heading_level_2 + 复核标记，不再强制 body_text
        assert result[1]["category"] == "heading_level_2"
        assert result[1]["needs_review"] is True
        assert "建议人工复核" in result[1]["comment"]

    def test_parse_batch_failure_fallback_to_single(self, temp_docx):
        """测试批量推理失败时降级到单条推理"""
        mock_single_result = {"label": "body_text", "score": 0.9}

        with patch(
            "wordformat.base.onnx_batch_infer", side_effect=RuntimeError("ONNX error")
        ):
            with patch(
                "wordformat.base.onnx_single_infer", return_value=mock_single_result
            ):
                base = DocxBase(temp_docx, "/fake/config.yaml")
                result = base.parse()

        # temp_docx 只有一个段落
        assert len(result) >= 1
        assert all(item["category"] == "body_text" for item in result)

    def test_parse_empty_document(self, tmp_path):
        """测试空文档的解析"""
        doc = Document()
        path = str(tmp_path / "empty.docx")
        doc.save(path)

        with patch("wordformat.base.onnx_batch_infer", return_value=[]) as mock_infer:
            base = DocxBase(path, "/fake/config.yaml")
            result = base.parse()

        assert result == []
        mock_infer.assert_not_called()

    def test_parse_batches_correctly(self, tmp_path):
        """测试按 BATCH_SIZE 分批推理"""
        path = self._create_multi_para_docx(tmp_path, [f"段落 {i}" for i in range(5)])

        call_count = 0

        def mock_batch(texts):
            nonlocal call_count
            call_count += 1
            return [{"label": "body_text", "score": 0.9}] * len(texts)

        with patch("wordformat.base.onnx_batch_infer", side_effect=mock_batch):
            with patch("wordformat.base.BATCH_SIZE", 2):
                base = DocxBase(path, "/fake/config.yaml")
                result = base.parse()

        assert len(result) == 5
        # BATCH_SIZE=2, 5 个段落 → 3 次调用 (2+2+1)
        assert call_count == 3

    def test_parse_includes_numbering_text(self, tmp_path):
        """测试解析时包含自动编号文字"""
        from docx.oxml import OxmlElement
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.parts.numbering import NumberingPart

        doc = Document()

        # 先移除已有的 numbering 关系
        rels = doc.part.rels
        to_remove = [k for k, v in rels.items() if v.reltype == RT.NUMBERING]
        for k in to_remove:
            del rels[k]

        p = doc.add_paragraph("绪论")

        # 添加 numbering
        numbering_elm = OxmlElement("w:numbering")
        abstract_num = OxmlElement("w:abstractNum")
        abstract_num.set(qn("w:abstractNumId"), "0")
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        numFmt = OxmlElement("w:numFmt")
        numFmt.set(qn("w:val"), "decimal")
        lvl.append(numFmt)
        lvlText = OxmlElement("w:lvlText")
        lvlText.set(qn("w:val"), "%1.")
        lvl.append(lvlText)
        abstract_num.append(lvl)
        numbering_elm.append(abstract_num)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), "1")
        abstract_num_id_ref = OxmlElement("w:abstractNumId")
        abstract_num_id_ref.set(qn("w:val"), "0")
        num.append(abstract_num_id_ref)
        numbering_elm.append(num)

        numbering_part = NumberingPart(
            PackURI("/word/numbering.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml",
            numbering_elm,
            doc.part.package,
        )
        doc.part.relate_to(numbering_part, RT.NUMBERING)

        pPr = OxmlElement("w:pPr")
        numPr = OxmlElement("w:numPr")
        numId_elem = OxmlElement("w:numId")
        numId_elem.set(qn("w:val"), "1")
        numPr.append(numId_elem)
        ilvl_elem = OxmlElement("w:ilvl")
        ilvl_elem.set(qn("w:val"), "0")
        numPr.append(ilvl_elem)
        pPr.append(numPr)
        p._element.insert(0, pPr)

        path = str(tmp_path / "numbered.docx")
        doc.save(path)

        with patch(
            "wordformat.base.onnx_batch_infer",
            return_value=[
                {"text": "1. 绪论", "label": "body_text", "pred_id": 0, "score": 0.9}
            ],
        ):
            base = DocxBase(path, "/fake/config.yaml")
            result = base.parse()

        assert len(result) == 1
        # 段落文本应包含编号 "1. 绪论"
        assert result[0]["paragraph"] == "1. 绪论"


# ============================================================
# base.py — 后处理链单元测试（C2–C6）
# ============================================================


def _mk_item(category, paragraph="", score=1.0):
    """构造一个 parse 输出风格的段落字典。"""
    return {
        "category": category,
        "paragraph": paragraph,
        "score": score,
        "comment": "",
        "needs_review": False,
    }


class TestPostProcessing:
    """直接测试 base.py 中的后处理函数。"""

    def test_apply_threshold_low_conf_keeps_label_and_flags(self):
        item = _mk_item("heading_level_2", "研究背景", score=0.3)
        _apply_threshold(item)
        assert item["category"] == "heading_level_2"  # 不再硬砍为 body_text
        assert item["needs_review"] is True
        assert "建议人工复核" in item["comment"]

    def test_apply_threshold_relaxed_for_abstract_en_title(self):
        # abstract_english_title 阈值 0.3，0.41 应通过（对应文档 Abstract 0.41 误伤）
        item = _mk_item("abstract_english_title", "Abstract", score=0.41)
        _apply_threshold(item)
        assert item["needs_review"] is False

    def test_apply_threshold_high_conf_ok(self):
        item = _mk_item("body_text", "正文", score=0.9)
        _apply_threshold(item)
        assert item["needs_review"] is False

    def test_fix_document_title_by_position(self):
        result = [
            _mk_item("body_text", "校园二手物品交易平台的设计与实现"),
            _mk_item("body_text", "摘要"),
            _mk_item("body_text", "本文研究了……"),
        ]
        _fix_document_title(result)
        assert result[0]["category"] == "document_title"

    def test_fix_document_title_skips_punctuated_first_para(self):
        result = [
            _mk_item("body_text", "这是一个很长的开头句子。"),
            _mk_item("body_text", "摘要"),
        ]
        _fix_document_title(result)
        assert result[0]["category"] == "body_text"  # 有句末标点 → 不改

    def test_fix_abstract_en_title_override(self):
        result = [_mk_item("body_text", "Abstract", score=0.41)]
        _fix_abstract_en_title(result)
        assert result[0]["category"] == "abstract_english_title"
        assert result[0]["score"] == 1.0

    def test_apply_footer_ai_notice(self):
        result = [_mk_item("body_text", "（注：文档部分内容可能由 AI 生成）")]
        _apply_footer(result)
        assert result[0]["category"] == "footer"

    def test_section_state_abstract_content_and_exit(self):
        result = [
            _mk_item("abstract_chinese_title", "摘要"),
            _mk_item("body_text", "本文研究电力大数据预测。"),
            _mk_item("body_text", "第二段摘要内容。"),
            _mk_item("keywords_chinese", "关键词：电力；大数据"),
            _mk_item("body_text", "关键词之后的普通正文"),
        ]
        _apply_section_state(result)
        assert result[1]["category"] == "abstract_chinese_content"
        assert result[2]["category"] == "abstract_chinese_content"
        # 关键词退出章节，后续 body_text 不再被吸收
        assert result[4]["category"] == "body_text"

    def test_section_state_references_requires_number(self):
        result = [
            _mk_item("references_title", "参考文献"),
            _mk_item("body_text", "[1] 张三. 论文题目[J]. 期刊, 2020."),
            _mk_item("body_text", "这一行没有编号，应保持 body_text"),
        ]
        _apply_section_state(result)
        assert result[1]["category"] == "references_content"
        assert result[2]["category"] == "body_text"

    def test_section_state_acknowledgements(self):
        result = [
            _mk_item("acknowledgements_title", "致谢"),
            _mk_item("body_text", "感谢导师的悉心指导。"),
        ]
        _apply_section_state(result)
        assert result[1]["category"] == "acknowledgements_content"

    def test_fix_known_categories_preserves_document_title(self):
        result = [
            _mk_item("document_title", "论文题目"),
            _mk_item("body_text", "学校：XX大学"),  # 封面 → other
            _mk_item("abstract_chinese_title", "摘要"),
            _mk_item("body_text", "本文……"),
        ]
        _fix_known_categories(result)
        assert result[0]["category"] == "document_title"  # 保留，不被覆盖为 other
        assert result[1]["category"] == "other"  # 封面跳过格式化
        assert result[2]["category"] == "abstract_chinese_title"

    def test_fix_sequence_references_rule5(self):
        result = [
            _mk_item("body_text", "参考文献"),  # 规则2 → references_title
            _mk_item("body_text", "[1] 张三. 论文题目[J]. 期刊, 2020."),
            _mk_item("body_text", "[2] 李四. 研究报告[R]. 2021."),
        ]
        _fix_sequence(result)
        assert result[0]["category"] == "references_title"
        assert result[1]["category"] == "references_content"
        assert result[2]["category"] == "references_content"

    def test_section_state_low_conf_title_does_not_propagate(self):
        """门控：低置信章节标题（score<0.6）不触发内容标签传播。

        源码附录里 "class X:"→abstract_english_title(0.45) 这类伪标题，
        不应把后续 body_text 放大成 abstract_english_content。
        """
        result = [
            _mk_item("abstract_english_title", "class XpathTool:", score=0.45),
            _mk_item("body_text", "def extract(self):"),
            _mk_item("body_text", "return data"),
        ]
        _apply_section_state(result)
        assert result[1]["category"] == "body_text"
        assert result[2]["category"] == "body_text"

    def test_section_state_high_conf_title_still_propagates(self):
        """门控：可信章节标题（score>=0.6）正常传播（无回归）。"""
        result = [
            _mk_item("references_title", "参考文献", score=0.96),
            _mk_item("body_text", "[1] 张三. 论文[J]. 期刊, 2020."),
        ]
        _apply_section_state(result)
        assert result[1]["category"] == "references_content"

    def test_neutralize_appendix_resets_frontmatter_labels(self):
        """附录区中性化：可信 heading_fulu 之后的摘要/关键词/致谢类→body_text。"""
        result = [
            _mk_item("heading_fulu", "附 录 B", score=0.97),
            _mk_item("abstract_english_title", "import json", score=0.52),
            _mk_item("abstract_english_content", "from pathlib import Path", 1.0),
            _mk_item("keywords_english", "def run(self):", score=0.96),
            _mk_item("acknowledgements_title", "return", score=0.41),
            _mk_item("figure_image", ""),
        ]
        _neutralize_appendix(result)
        assert result[1]["category"] == "body_text"
        assert result[2]["category"] == "body_text"
        assert result[3]["category"] == "body_text"
        assert result[4]["category"] == "body_text"
        assert result[5]["category"] == "figure_image"  # 图片不受影响

    def test_neutralize_appendix_keeps_frontmatter_before_appendix(self):
        """附录之前（正文/致谢区）不受中性化影响。"""
        result = [
            _mk_item("acknowledgements_title", "致 谢", score=0.98),
            _mk_item("keywords_english", "Keywords: A; B", score=0.98),
            _mk_item("heading_fulu", "附 录 A", score=0.97),
        ]
        _neutralize_appendix(result)
        assert result[0]["category"] == "acknowledgements_title"
        assert result[1]["category"] == "keywords_english"

    def test_neutralize_appendix_low_conf_fulu_does_not_enter(self):
        """低置信 heading_fulu 不触发附录区，避免误判扩大中性化范围。"""
        result = [
            _mk_item("heading_fulu", "附录", score=0.3),
            _mk_item("keywords_english", "Keywords: A; B", score=0.98),
        ]
        _neutralize_appendix(result)
        assert result[1]["category"] == "keywords_english"

    def test_neutralize_appendix_exits_on_confident_heading_level_1(self):
        """可信 heading_level_1 结束附录区，后续真实标签保留。"""
        result = [
            _mk_item("heading_fulu", "附 录 A", score=0.97),
            _mk_item("keywords_english", "code line", score=0.9),
            _mk_item("heading_level_1", "7 结论", score=0.95),
            _mk_item("keywords_english", "Keywords: real", score=0.98),
        ]
        _neutralize_appendix(result)
        assert result[1]["category"] == "body_text"  # 附录内→中性化
        assert result[3]["category"] == "keywords_english"  # 已退出附录，保留

    def test_fix_toc_overrides_to_heading_mulu(self):
        """目录：独立成段的“目录”→heading_mulu（模型词表无此标签）。"""
        result = [
            _mk_item("heading_fulu", "目录", score=0.49),
            _mk_item("body_text", "目 录"),
        ]
        _fix_toc(result)
        assert result[0]["category"] == "heading_mulu"
        assert result[1]["category"] == "heading_mulu"  # 带空格也匹配

    def test_fix_toc_ignores_non_toc_text(self):
        """非“目录”整段文本不被改写。"""
        result = [_mk_item("body_text", "见目录第 3 页")]
        _fix_toc(result)
        assert result[0]["category"] == "body_text"

    def test_fix_abstract_titles_promotes_cn_title(self):
        """摘要页中文标题：紧邻“摘要：”上方的 other 标题行 → abstract_chinese_title。"""
        result = [
            _mk_item("other", "导师签名：日期："),  # 封面签名，不应升格
            _mk_item("other", "基于单片机的公交车汉字显示系统"),  # CN 标题
            _mk_item("abstract_chinese_title_content", "摘要：本文研究了……"),
        ]
        _fix_abstract_titles(result)
        assert result[1]["category"] == "abstract_chinese_title"
        assert result[1]["score"] == 1.0
        assert result[0]["category"] == "other"  # 非最近邻的签名行不受影响

    def test_fix_abstract_titles_promotes_en_title(self):
        """摘要页英文标题：紧邻“Abstract:”上方被误判为 keywords_english 的标题 → abstract_english_title。"""
        result = [
            _mk_item("keywords_chinese", "关键词：汉字显示；单片机"),
            _mk_item(
                "keywords_english",
                "Single-chip Microcontroller-Based Chinese Character Display System",
                score=0.28,
            ),
            _mk_item("abstract_english_title_content", "Abstract:With the acceleration of……"),
        ]
        _fix_abstract_titles(result)
        assert result[1]["category"] == "abstract_english_title"
        assert result[1]["needs_review"] is False

    def test_fix_abstract_titles_skips_cover_signature(self):
        """无独立标题时，签名/日期行（冒号结尾或含封面特征词）不被误升为摘要标题。"""
        result = [
            _mk_item("other", "导师签名：日期："),
            _mk_item("abstract_chinese_title_content", "摘要：本文研究了……"),
        ]
        _fix_abstract_titles(result)
        assert result[0]["category"] == "other"

    def test_fix_abstract_titles_skips_long_sentence(self):
        """上方是长句（句末标点）时不升格，避免把正文误判为标题。"""
        result = [
            _mk_item("body_text", "这是一段很长的正文内容，讲述了研究背景。"),
            _mk_item("abstract_chinese_title_content", "摘要：本文研究了……"),
        ]
        _fix_abstract_titles(result)
        assert result[0]["category"] == "body_text"

    def test_fix_abstract_titles_skips_already_title(self):
        """上方已是摘要标题时不重复处理（无回归）。"""
        result = [
            _mk_item("abstract_chinese_title", "校园二手交易平台的设计与实现"),
            _mk_item("abstract_chinese_title_content", "摘要：本文研究了……"),
        ]
        _fix_abstract_titles(result)
        assert result[0]["category"] == "abstract_chinese_title"

    def test_fix_sequence_acknowledgements_rule6(self):
        """规则6：致谢标题（规则3 补上）之后的 body_text → acknowledgements_content。"""
        result = [
            _mk_item("body_text", "致  谢"),  # 规则3 → acknowledgements_title
            _mk_item("body_text", "在四年本科的学习生涯中，我感谢导师。"),
            _mk_item("body_text", ""),  # 空段保持 body_text
        ]
        _fix_sequence(result)
        assert result[0]["category"] == "acknowledgements_title"
        assert result[1]["category"] == "acknowledgements_content"
        assert result[2]["category"] == "body_text"

    def test_fix_sequence_acknowledgements_rule6_gated(self):
        """规则6 门控：低置信 acknowledgements_title（如源码 "{"@0.39）不把代码扫成致谢正文。"""
        result = [
            _mk_item("body_text", "void Monitor_function(void)"),
            _mk_item("acknowledgements_title", "{", score=0.39),  # 模型误判的低置信标题
            _mk_item("body_text", "    if(USART2_WaitRecive() == 0)"),
            _mk_item("body_text", "    TIM_SetCompare1(TIM2, 1900);"),
        ]
        _fix_sequence(result)
        assert result[2]["category"] == "body_text"  # 门控生效，代码不被扫成致谢正文
        assert result[3]["category"] == "body_text"


# ============================================================
# utils.py — _format_number 额外覆盖测试
# ============================================================
