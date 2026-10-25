#! /usr/bin/env python
# @Time    : 2026/10/24
# @Author  : afish
# @File    : test_gj_math.py
"""gj 公式段识别：OMML 结构短路判定为 math_block（模型词表无公式类）。"""

from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from wordformat.base import DocxBase

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _add_math(p, text="E=mc2", para_level=False):
    omml = f'<m:oMath xmlns:m="{M}"><m:r><m:t>{text}</m:t></m:r></m:oMath>'
    if para_level:
        omml = f'<m:oMathPara xmlns:m="{M}">{omml}</m:oMathPara>'
    p._p.append(parse_xml(omml))


def _add_text(p, text):
    p._p.append(parse_xml(f'<w:r xmlns:w="{W}"><w:t>{text}</w:t></w:r>'))


def _parse(path, n_texts):
    """n_texts 必须等于真正送进 AI 的段数；少了会触发 parse() 的未处理断言。"""
    mock_results = [{"label": "body_text", "score": 0.9} for _ in range(n_texts)]
    with patch(
        "wordformat.base.onnx_batch_infer",
        return_value=(mock_results, "body_text"),
    ):
        return DocxBase(path, "/fake/config.yaml").parse()


class TestGjMathBlock:
    def test_empty_paragraph_with_omml_is_math_block(self, tmp_path):
        doc = Document()
        _add_math(doc.add_paragraph())
        path = tmp_path / "math1.docx"
        doc.save(path)

        result = _parse(str(path), n_texts=0)
        assert result[0]["category"] == "math_block"
        assert result[0]["comment"] == "公式段落"

    def test_display_formula_with_number_keeps_number_text(self, tmp_path):
        doc = Document()
        p = doc.add_paragraph()
        _add_math(p, para_level=True)
        _add_text(p, "（2-1）")
        path = tmp_path / "math2.docx"
        doc.save(path)

        result = _parse(str(path), n_texts=0)
        assert result[0]["category"] == "math_block"
        # 编号文字必须保留，否则 jd 重建会丢掉 (2-1)
        assert result[0]["paragraph"] == "（2-1）"

    def test_prose_with_inline_math_is_not_math_block(self, tmp_path):
        doc = Document()
        p = doc.add_paragraph()
        _add_text(p, "由式（2-1）可得，其中 x 为位移。")
        _add_math(p, "x2")
        path = tmp_path / "math3.docx"
        doc.save(path)

        # 送 AI 1 段：n_texts=0 会因未处理断言失败，可反证它没被短路
        result = _parse(str(path), n_texts=1)
        assert result[0]["category"] == "body_text"

    def test_number_without_math_still_goes_to_ai(self, tmp_path):
        doc = Document()
        doc.add_paragraph("（2-1）")
        path = tmp_path / "math4.docx"
        doc.save(path)

        result = _parse(str(path), n_texts=1)
        assert result[0]["category"] == "body_text"

    def test_equation_number_bracket_variants(self, tmp_path):
        doc = Document()
        for number in ("（2-1）", "(2.1)", "[3-12]", "\t（2-1）\t", "（12—3）"):
            p = doc.add_paragraph()
            _add_math(p)
            _add_text(p, number)
        path = tmp_path / "math5.docx"
        doc.save(path)

        result = _parse(str(path), n_texts=0)
        assert [r["category"] for r in result] == ["math_block"] * 5

    def test_prose_number_prefix_is_not_matched(self, tmp_path):
        """"由式（2-1）"这类带前后文的编号不短路，避免吞掉正文。"""
        doc = Document()
        for text in ("式（2-1）成立", "（2-1）其中 x 为位移", "（2-1-3）"):
            p = doc.add_paragraph()
            _add_math(p)
            _add_text(p, text)
        path = tmp_path / "math6.docx"
        doc.save(path)

        result = _parse(str(path), n_texts=3)
        assert [r["category"] for r in result] == ["body_text"] * 3

    def test_image_paragraph_branch_unaffected(self, tmp_path):
        doc = Document()
        doc.add_paragraph("")
        path = tmp_path / "image.docx"
        doc.save(path)

        with patch("wordformat.base.para_contains_image", return_value=True):
            result = _parse(str(path), n_texts=0)
        assert result[0]["category"] == "figure_image"

    def test_blank_paragraph_still_body_text(self, tmp_path):
        doc = Document()
        doc.add_paragraph("")
        path = tmp_path / "blank.docx"
        doc.save(path)

        result = _parse(str(path), n_texts=0)
        assert result[0]["category"] == "body_text"
        assert result[0]["comment"] == "空段落"


class TestMathBlockFormatting:
    """识别出来后要能真的排版：纯公式段没有 w:r，段落格式仍须写入 pPr。"""

    def _math_para(self, with_number=False):
        doc = Document()
        p = doc.add_paragraph()
        _add_math(p, para_level=True)
        if with_number:
            _add_text(p, "（2-1）")
        return doc, p

    def test_apply_centers_runless_formula_paragraph(self):
        from unittest.mock import patch as _patch

        from wordformat.config.models import NodeConfigRoot
        from wordformat.rules.math import MathBlock

        doc, p = self._math_para()
        assert len(p.runs) == 0
        node = MathBlock(value="", level=5, paragraph=p)
        node.load_config(NodeConfigRoot())
        with _patch.object(node, "add_comment"):
            node.apply_format(doc)
        jc = p._p.pPr.find(qn("w:jc"))
        assert jc.get(qn("w:val")) == "center"

    def test_check_mode_skips_runless_paragraph_without_crash(self):
        """check 模式需要 run 锚定批注，无 run 时静默跳过而不是报错。"""
        from wordformat.config.models import NodeConfigRoot
        from wordformat.rules.math import MathBlock

        doc, p = self._math_para()
        node = MathBlock(value="", level=5, paragraph=p)
        node.load_config(NodeConfigRoot())
        node.check_format(doc)
        assert p._p.pPr is None or p._p.pPr.find(qn("w:jc")) is None

    def test_plain_empty_paragraph_still_skipped(self):
        """无 run 且无公式的空段仍跳过：不给所有占位空段写 pPr。"""
        from wordformat.config.models import NodeConfigRoot
        from wordformat.rules.body import BodyText

        doc = Document()
        p = doc.add_paragraph("")
        node = BodyText(value="", level=5, paragraph=p)
        node.load_config(NodeConfigRoot())
        node.apply_format(doc)
        jc = None if p._p.pPr is None else p._p.pPr.find(qn("w:jc"))
        assert jc is None

    def test_formula_with_number_still_reports_and_applies(self):
        from unittest.mock import patch as _patch

        from wordformat.config.models import NodeConfigRoot
        from wordformat.rules.math import MathBlock

        doc, p = self._math_para(with_number=True)
        node = MathBlock(value="（2-1）", level=5, paragraph=p)
        node.load_config(NodeConfigRoot())
        with _patch.object(node, "add_comment") as comment:
            node.check_format(doc)
            node.apply_format(doc)
        jc = p._p.pPr.find(qn("w:jc"))
        assert jc.get(qn("w:val")) == "center"
        # 有 run 时可正常锚定批注（apply 前的 check 报出不对齐）
        assert comment.called
